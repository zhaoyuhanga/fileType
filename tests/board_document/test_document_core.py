"""墨软文档：格式矩阵 / 解析 / 写出 / 美化 / 质量 / 存储（引擎层测试）。"""
from __future__ import annotations

import base64
import json
import zipfile
from pathlib import Path

import pytest

from modu_workbench.core.document import (
    CATEGORY_ORDER,
    EXPORT_TARGETS,
    FORMAT_SPECS,
    BeautifyOptions,
    DocumentIR,
    DocumentLibrary,
    DocumentStorage,
    analyze_document,
    beautify_document,
    describe_matrix,
    format_key_for_path,
    format_label,
    get_spec,
    heuristic_score,
    parse_document,
    parse_html,
    parse_markdown,
    parse_plain_text,
    template_choices,
    template_options,
    write_document,
)
from modu_workbench.core.document import formats as fmt


# ---------------------------------------------------------------- 格式矩阵


def test_format_matrix_is_consistent() -> None:
    keys = [spec.key for spec in FORMAT_SPECS]
    assert len(keys) == len(set(keys)), "格式 key 必须唯一"
    for spec in FORMAT_SPECS:
        assert spec.category in CATEGORY_ORDER, f"{spec.key} 的分类未登记：{spec.category}"
        assert spec.extensions, f"{spec.key} 没有扩展名"
        assert spec.label and spec.engine
        assert spec.capability_text()


def test_format_lookup_by_path_and_extension() -> None:
    assert format_key_for_path("a/b/报告.docx") == "docx"
    assert format_key_for_path("说明.MD") == "markdown"
    assert format_key_for_path("数据.yml") == "yaml"
    assert format_key_for_path("扫描件.jpeg") == "jpg"
    assert format_key_for_path("宏.docm") == "docx"
    assert format_key_for_path("脚本.py") == "code"
    assert format_key_for_path("未知.zzz") == ""
    assert format_label("markdown") == "Markdown"


def test_capability_matrix_is_honest() -> None:
    """只读格式不能声称可编辑，OCR 只出现在能识别的格式上。"""
    assert get_spec("docx").edit is True
    assert get_spec("xlsx").edit is True
    assert get_spec("pdf").edit is False
    assert get_spec("pdf").ocr is True
    assert get_spec("doc").edit is False
    assert get_spec("xls").edit is False
    assert get_spec("png").render == fmt.RENDER_IMAGE
    assert len(fmt.editable_formats()) < len(FORMAT_SPECS)


def test_export_targets_have_extensions() -> None:
    for key, label in EXPORT_TARGETS:
        assert key in fmt.EXPORT_EXTENSIONS, f"{key} 没有扩展名映射"
        assert label
    assert describe_matrix()[0]["label"] == FORMAT_SPECS[0].label


def test_extension_dialog_filter() -> None:
    text = fmt.extensions_for_dialog()
    assert "*.docx" in text and "所有文件" in text


# ---------------------------------------------------------------- 解析


def test_parse_markdown_blocks() -> None:
    ir = parse_markdown(
        "# 标题\n\n正文一。\n\n## 1.1 小节\n\n- 项目一\n- 项目二\n\n"
        "| A | B |\n|---|---|\n| 1 | 2 |\n\n> 引用\n\n```python\nprint(1)\n```\n\n---\n",
        title="示例")
    kinds = [block.kind for block in ir.blocks]
    assert "heading" in kinds and "table" in kinds and "code" in kinds
    assert "list_item" in kinds and "quote" in kinds and "page_break" in kinds
    assert ir.outline()[0] == (1, "标题", 0)
    assert ir.stats()["tables"] == 1
    assert "# 标题" in ir.markdown()


def test_parse_plain_text_detects_headings() -> None:
    ir = parse_plain_text("第一章 概述\n\n这是正文内容，长度足够。\n\n一、小节标题\n\n更多正文。")
    headings = [block.text for block in ir.headings()]
    assert headings == ["第一章 概述", "一、小节标题"]
    assert ir.blocks[1].kind == "paragraph"


def test_parse_html_blocks() -> None:
    ir = parse_html("<html><body><h1>标题</h1><p>正文</p><table>"
                    "<tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr>"
                    "</table></body></html>")
    assert [block.kind for block in ir.blocks] == ["heading", "paragraph", "table"]
    assert ir.tables()[0].rows[1] == ["1", "2"]


def test_parse_csv_and_json(tmp_path: Path) -> None:
    csv_file = tmp_path / "data.csv"
    csv_file.write_text("名称,数量\nA,1\nB,2\n", encoding="utf-8")
    ir = parse_document(csv_file)
    assert ir.first_table().rows[0] == ["名称", "数量"]
    assert ir.metadata["format"] == "csv"

    json_file = tmp_path / "data.json"
    json_file.write_text('{"a": 1, "b": [2, 3]}', encoding="utf-8")
    ir = parse_document(json_file)
    assert '"a": 1' in ir.blocks[0].text
    assert "2 个键" in ir.metadata["structure"]


def test_parse_json_error_is_readable(tmp_path: Path) -> None:
    bad = tmp_path / "bad.json"
    bad.write_text('{"a": 1,}', encoding="utf-8")
    with pytest.raises(Exception) as info:
        parse_document(bad)
    assert "JSON 语法错误" in str(info.value)


def test_parse_unsupported_extension(tmp_path: Path) -> None:
    weird = tmp_path / "a.zzz"
    weird.write_text("x", encoding="utf-8")
    with pytest.raises(Exception) as info:
        parse_document(weird)
    assert "暂不支持" in str(info.value)


def test_parse_docx_roundtrip(tmp_path: Path) -> None:
    docx_file = tmp_path / "报告.docx"
    source = parse_markdown("# 报告\n\n## 1.1 数据\n\n正文段落。\n\n"
                            "| A | B |\n|---|---|\n| 1 | 2 |\n", title="报告")
    write_document(source, docx_file, "docx")
    assert docx_file.is_file()

    parsed = parse_document(docx_file)
    assert parsed.format_key == "docx"
    assert [block.text for block in parsed.headings()] == ["报告", "1.1 数据"]
    assert parsed.tables()[0].rows[1] == ["1", "2"]


def test_parse_xlsx_keeps_formulas(tmp_path: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "明细"
    sheet.append(["项目", "数量", "单价", "金额"])
    sheet.append(["A", 3, 10, "=B2*C2"])
    workbook.save(tmp_path / "t.xlsx")

    ir = parse_document(tmp_path / "t.xlsx")
    table = ir.first_table()
    assert table.name == "明细"
    assert table.rows[1][3] == "=B2*C2"
    assert table.meta["formula_cells"] == 1


def test_parse_missing_file(tmp_path: Path) -> None:
    with pytest.raises(Exception) as info:
        parse_document(tmp_path / "nope.md")
    assert "不存在" in str(info.value)


# ---------------------------------------------------------------- 图片素材（Word）

#: 1×1 的合法 PNG（不依赖 Pillow，测试里当"插图"用）
PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8DwHwAFAAH/q842iQAAAABJRU5ErkJggg==")


def _docx_with_image(path: Path) -> Path:
    from docx import Document
    from docx.shared import Inches

    image = path.parent / "插图.png"
    image.write_bytes(PNG_1PX)
    document = Document()
    document.add_heading("带图的文档", level=1)
    document.add_paragraph("图片前正文。")
    document.add_picture(str(image), width=Inches(1.0))
    document.add_paragraph("图片后正文。")
    document.save(path)
    return path


def _media_parts(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as bundle:
        return [name for name in bundle.namelist() if name.startswith("word/media/")]


def test_docx_images_are_parsed_into_blocks(tmp_path: Path) -> None:
    """图片要变成 image 块（过去直接丢掉：段落没文字就 continue）。"""
    source = _docx_with_image(tmp_path / "带图.docx")
    assert _media_parts(source), "测试样本本身要带图"
    ir = parse_document(source)
    images = [block for block in ir.blocks if block.kind == "image"]
    assert len(images) == 1
    info = images[0].meta["image"]
    assert info["file"].endswith(".png") and info["bytes"] == len(PNG_1PX)
    # 图片字节不进 IR：版本快照只存引用，否则一份带图文档的 40 个快照会吃掉几十 MB
    payload = json.dumps(ir.to_dict(), ensure_ascii=False)
    assert "base64" not in payload and len(payload) < 4096
    assert [block.kind for block in ir.blocks] == ["heading", "paragraph", "image", "paragraph"]


def test_docx_images_survive_save_and_export(tmp_path: Path) -> None:
    """打开带图 Word → 保存/导出后图片还在（原来是静默丢图）。"""
    source = _docx_with_image(tmp_path / "带图.docx")
    ir = parse_document(source)

    library = DocumentLibrary(DocumentStorage(tmp_path / "m.db"), output_dir=tmp_path / "out")
    library.save(ir, source)
    assert _media_parts(source), "保存回原文件后图片必须还在"

    reparsed = parse_document(source)
    assert any(block.kind == "image" for block in reparsed.blocks)

    exported = write_document(ir, tmp_path / "导出.docx", "docx")
    assert _media_parts(exported.path), "导出 Word 也要带图"
    assert exported.notes == []


def test_docx_image_missing_cache_keeps_placeholder(tmp_path: Path) -> None:
    """素材丢了也不能静默：留一行占位文字并给出提示。"""
    from modu_workbench.core.document import media

    source = _docx_with_image(tmp_path / "带图.docx")
    ir = parse_document(source)
    info = next(block for block in ir.blocks if block.kind == "image").meta["image"]
    (media.media_dir() / info["file"]).unlink()

    result = write_document(ir, tmp_path / "无素材.docx", "docx")
    assert result.notes and "图片" in result.notes[0]
    text = parse_document(result.path).text()
    assert "[图片]" in text


def test_docx_image_in_other_targets(tmp_path: Path) -> None:
    """HTML 内嵌（自包含）、纯文本/Markdown 给占位说明，并明确提示带不了图。"""
    source = _docx_with_image(tmp_path / "带图.docx")
    ir = parse_document(source)
    name = next(block for block in ir.blocks if block.kind == "image").text
    assert name, "图片块要有展示名"

    html = write_document(ir, tmp_path / "a.html", "html")
    assert "data:image/png;base64," in html.path.read_text(encoding="utf-8")

    text = write_document(ir, tmp_path / "a.txt", "txt")
    assert f"[图片] {name}" in text.path.read_text(encoding="utf-8")
    assert text.notes and "不能内嵌图片" in text.notes[0]

    markdown = write_document(ir, tmp_path / "a.md", "markdown")
    assert f"_[图片：{name}]_" in markdown.path.read_text(encoding="utf-8")


def test_beautify_keeps_images(tmp_path: Path) -> None:
    """美化（编号/样式统一/目录）不能把图片块吃掉。"""
    source = _docx_with_image(tmp_path / "带图.docx")
    ir = parse_document(source)
    result = beautify_document(ir, template="report")
    assert any(block.kind == "image" for block in result.ir.blocks)
    exported = write_document(result.ir, tmp_path / "美化后.docx", "docx")
    assert _media_parts(exported.path)


def test_markdown_images_become_blocks(tmp_path: Path) -> None:
    """Markdown 里的 `![]()`：本地相对路径收进素材缓存，远程 URL 原样保留。"""
    picture = tmp_path / "插图.png"
    picture.write_bytes(PNG_1PX)
    source = tmp_path / "带图.md"
    source.write_text(
        "# 报告\n\n![本地图](插图.png)\n\n![远程图](https://example.com/a.png)\n",
        encoding="utf-8")

    ir = parse_document(source)
    images = [block for block in ir.blocks if block.kind == "image"]
    assert len(images) == 2
    local, remote = images
    assert local.meta["image"]["bytes"] == len(PNG_1PX)          # 本地图进了素材缓存
    assert remote.meta["src"] == "https://example.com/a.png"     # 远程图不抓取，只记地址

    # 注释：本地图导出 HTML 自包含；远程图照写 URL；Markdown 里远程图写成真实图片语法
    html = write_document(ir, tmp_path / "a.html", "html")
    body = html.path.read_text(encoding="utf-8").split("<body>", 1)[1]
    assert "data:image/png;base64," in body
    assert "https://example.com/a.png" in body
    markdown = write_document(ir, tmp_path / "a.md", "markdown")
    text = markdown.path.read_text(encoding="utf-8")
    assert "![远程图](https://example.com/a.png)" in text       # 有地址的写成真图片语法
    assert "_[图片：插图.png]_" in text                          # 缓存里的只留说明


def test_html_data_uri_image_is_kept(tmp_path: Path) -> None:
    """HTML 里内联的 base64 图片（保存网页很常见）也要收进素材缓存，而不是丢掉。"""
    encoded = base64.b64encode(PNG_1PX).decode("ascii")
    source = tmp_path / "网页.html"
    source.write_text(
        f"<html><body><h1>标题</h1><p>正文</p>"
        f"<img alt='内联图' src='data:image/png;base64,{encoded}'></body></html>",
        encoding="utf-8")

    ir = parse_document(source)
    image = next(block for block in ir.blocks if block.kind == "image")
    assert image.meta["image"]["bytes"] == len(PNG_1PX)
    exported = write_document(ir, tmp_path / "导出.docx", "docx")
    assert _media_parts(exported.path), "内联图片也要能写进 Word"


def test_docx_textbox_and_revision_text_is_not_lost(tmp_path: Path) -> None:
    """文本框与修订插入里的文字不能因为 python-docx 的 paragraph.text 而消失。"""
    import docx as docx_lib
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    source = tmp_path / "修订.docx"
    document = docx_lib.Document()
    document.add_paragraph("正常段落。")
    paragraph = document.add_paragraph()
    inserted = OxmlElement("w:ins")
    inserted.set(qn("w:id"), "1")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "修订新增的一句。"
    run.append(text)
    inserted.append(run)
    paragraph._p.append(inserted)
    deleted = OxmlElement("w:del")
    deleted.set(qn("w:id"), "2")
    deleted_run = OxmlElement("w:r")
    deleted_text = OxmlElement("w:delText")
    deleted_text.text = "这句已经被删掉了。"
    deleted_run.append(deleted_text)
    deleted.append(deleted_run)
    paragraph._p.append(deleted)

    box_paragraph = document.add_paragraph()
    # 文本框是 VML（`v:` 前缀不在 python-docx 的 nsmap 里），直接按 OOXML 片段建
    from docx.oxml import parse_xml

    box_paragraph._p.append(parse_xml(
        '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        ' xmlns:v="urn:schemas-microsoft-com:vml">'
        '<v:shape style="width:200pt;height:40pt"><v:textbox><w:txbxContent>'
        "<w:p><w:r><w:t>文本框里的说明</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict>"))
    document.save(source)

    ir = parse_document(source)
    text_all = "\n".join(block.text for block in ir.blocks)
    assert "修订新增的一句。" in text_all, "修订插入的文字不能丢"
    assert "这句已经被删掉了。" not in text_all, "被删除的修订内容不该出现在正文里"
    assert "文本框里的说明" in text_all, "文本框文字要退化为引用块而不是丢掉"
    assert any(block.style.get("textbox") for block in ir.blocks)
    assert any("修订" in warning for warning in ir.warnings)


def test_filename_title_is_not_invented_as_visible_title(tmp_path: Path) -> None:
    """标题只是文件名时不要造一行可见标题：否则"打开 Word → 保存"会平白多一行。"""
    source = tmp_path / "年度总结.docx"
    document_with_body = _docx_with_image(source)
    ir = parse_document(document_with_body)
    assert ir.metadata.get("title_source") == "filename"
    exported = write_document(ir, tmp_path / "导出.docx", "docx")
    text = parse_document(exported.path).text()
    assert "年度总结" not in text, "文件名不该变成正文标题"

    # 显式给了标题（新建文档/AI 起标题）时照旧写出来
    explicit = parse_markdown("# 概述\n\n正文。", title="年度报告")
    html = write_document(explicit, tmp_path / "b.html", "html").path.read_text(encoding="utf-8")
    assert "年度报告" in html and "概述" in html


def test_media_cache_stats_and_clear(tmp_path: Path) -> None:
    """素材缓存可以看大小、可以清空（清空后图片变占位，重新解析会补回来）。"""
    from modu_workbench.core.document import media

    source = _docx_with_image(tmp_path / "带图.docx")
    parse_document(source)
    stats = media.cache_stats()
    assert stats["files"] >= 1 and stats["bytes"] > 0

    removed = media.clear_cache()
    assert removed["files"] == stats["files"]
    assert media.cache_stats()["files"] == 0

    again = parse_document(source)
    assert any(block.kind == "image" for block in again.blocks)
    assert media.cache_stats()["files"] >= 1


# ---------------------------------------------------------------- 写出


@pytest.mark.parametrize("target", ["txt", "markdown", "html", "docx", "xlsx", "csv", "tsv", "json"])
def test_write_targets(tmp_path: Path, target: str) -> None:
    ir = parse_markdown("# 标题\n\n正文。\n\n| A | B |\n|---|---|\n| 1 | 2 |\n", title="标题")
    result = write_document(ir, tmp_path / f"out.{target}", target)
    assert result.path.is_file() and result.path.stat().st_size > 0


def test_write_does_not_overwrite(tmp_path: Path) -> None:
    ir = DocumentIR.from_text("内容", title="文稿")
    first = write_document(ir, tmp_path / "a.txt", "txt")
    second = write_document(ir, tmp_path / "a.txt", "txt")
    assert first.path.name == "a.txt"
    assert second.path.name == "a (1).txt"


def test_write_json_payload_has_blocks(tmp_path: Path) -> None:
    ir = parse_markdown("# 标题\n\n正文。", title="标题")
    result = write_document(ir, tmp_path / "a.json", "json")
    payload = json.loads(result.path.read_text(encoding="utf-8"))
    assert payload["title"] == "标题"
    assert payload["stats"]["blocks"] >= 1
    assert payload["blocks"][0]["kind"] == "heading"


def test_write_html_keeps_brand_and_table(tmp_path: Path) -> None:
    ir = parse_markdown("| A | B |\n|---|---|\n| 1 | 2 |", title="表格稿")
    options = BeautifyOptions(brand_color="#B01F24", header_text="内部资料")
    result = write_document(ir, tmp_path / "a.html", "html", options=options)
    html = result.path.read_text(encoding="utf-8")
    assert "#B01F24" in html and "内部资料" in html and "<table>" in html


def test_write_does_not_repeat_document_title(tmp_path: Path) -> None:
    """标题与首个标题块同名时只输出一次。

    `我的报告.md` 解析出的 `ir.title` 就是首个 H1 的文字，美化也会用首个标题补
    `ir.title`（`beautify._document_title_block`）——再单独写一遍标题就是重复。
    """
    ir = parse_markdown("# 我的报告\n\n正文。", title="我的报告")
    html = write_document(ir, tmp_path / "a.html", "html").path.read_text(encoding="utf-8")
    body = html.split("<body>", 1)[1]          # <head> 里的 <title> 不算正文
    assert body.count("我的报告") == 1
    assert "<h1 class='title'>" not in body

    docx_path = tmp_path / "a.docx"
    write_document(ir, docx_path, "docx")
    assert parse_document(docx_path).text().count("我的报告") == 1


def test_write_keeps_title_when_it_differs_from_first_heading(tmp_path: Path) -> None:
    """标题与首个标题块不同名时照旧两行都写（不能把标题吃掉）。"""
    ir = parse_markdown("# 概述\n\n正文。", title="年度报告")
    html = write_document(ir, tmp_path / "a.html", "html").path.read_text(encoding="utf-8")
    assert "年度报告" in html and "概述" in html

    docx_path = tmp_path / "a.docx"
    write_document(ir, docx_path, "docx")
    text = parse_document(docx_path).text()
    assert "年度报告" in text and "概述" in text


def test_write_pdf_with_qapp(qapp, tmp_path: Path) -> None:  # noqa: ANN001
    ir = parse_plain_text("这是一段用于生成 PDF 的中文文本。", title="PDF 测试")
    result = write_document(ir, tmp_path / "a.pdf", "pdf")
    assert result.path.is_file() and result.path.stat().st_size > 1000


def test_write_xlsx_table_features(tmp_path: Path) -> None:
    """导出 xlsx 的表格能力：冻结表头、条件格式、数据验证。"""
    import openpyxl

    source = parse_markdown(
        "| 部门 | 状态 | 金额 |\n| --- | --- | --- |\n"
        "| 销售 | 完成 | 10 |\n| 研发 | 进行 | 20 |\n| 财务 | 完成 | 30 |\n",
        title="明细")
    options = BeautifyOptions(conditional_format=True, data_validation=True,
                              freeze_header=True)
    result = write_document(source, tmp_path / "t.xlsx", "xlsx", options=options)
    assert any("条件格式" in note for note in result.notes)
    assert any("数据验证" in note for note in result.notes)

    workbook = openpyxl.load_workbook(result.path)
    sheet = workbook.active
    assert sheet.freeze_panes == "A2"
    rules = list(sheet.conditional_formatting)
    assert rules and any("AVERAGE" in str(rule.rules[0].formula) for rule in rules)
    validations = list(sheet.data_validations.dataValidation)
    assert validations, "取值少的列应当写入下拉数据验证"
    formulas = " ".join(str(item.formula1) for item in validations)
    assert "完成" in formulas, "状态列应有下拉清单"
    assert '"销售,研发,财务"' in formulas, "部门列应有下拉清单"
    assert "10,20,30" not in formulas, "纯数值列不该塞下拉清单"
    workbook.close()


def test_write_xlsx_can_skip_table_features(tmp_path: Path) -> None:
    import openpyxl

    source = parse_markdown("| A | B |\n| --- | --- |\n| 1 | 2 |\n| 3 | 4 |\n| 5 | 6 |\n",
                            title="表")
    result = write_document(source, tmp_path / "plain.xlsx", "xlsx",
                           options=BeautifyOptions(conditional_format=False,
                                                   data_validation=False,
                                                   freeze_header=False))
    workbook = openpyxl.load_workbook(result.path)
    sheet = workbook.active
    assert sheet.freeze_panes is None
    assert list(sheet.conditional_formatting) == []
    assert list(sheet.data_validations.dataValidation) == []
    workbook.close()


# ---------------------------------------------------------------- 美化


def test_beautify_report_template() -> None:
    ir = parse_markdown("# 年度报告\n\n## 1.1 概述\n\n正文一。\n\n## 1.2 数据\n\n正文二。\n",
                        title="年度报告")
    result = beautify_document(ir, template="report")
    kinds = [change.kind for change in result.changes]
    assert "样式统一" in kinds and "生成目录" in kinds
    assert result.toc_entries >= 2
    # 文档标题不参与编号，下级标题从 1.x 开始
    assert result.ir.blocks[0].text == "年度报告"
    assert any(block.text.startswith("1.") for block in result.ir.headings())


def test_beautify_is_idempotent() -> None:
    ir = parse_markdown("# 标题\n\n## 1.1 小节\n\n正文。", title="标题")
    first = beautify_document(ir, template="report").ir
    second = beautify_document(first, template="report")
    assert second.ir.markdown() == first.markdown()
    assert second.change_count <= 1


def test_beautify_gov_template_styles() -> None:
    ir = parse_markdown("# 通知\n\n正文内容。", title="通知")
    options = template_options("gov")
    assert options.font_cn == "仿宋_GB2312"
    result = beautify_document(ir, options=options)
    paragraph = [block for block in result.ir.blocks if block.kind == "paragraph"][0]
    assert paragraph.style["font_cn"] == "仿宋_GB2312"
    assert paragraph.style["size"] == 16.0


def test_beautify_tables_flags() -> None:
    ir = parse_markdown("| A | B |\n|---|---|\n| 1 | 2 |", title="表")
    result = beautify_document(ir, template="report")
    table = result.ir.tables()[0]
    assert table.styles["border"] == "grid"
    assert table.styles["zebra"] is True
    assert len(table.styles["column_widths"]) == 2


def test_template_choices_cover_all_templates() -> None:
    keys = [key for key, _label in template_choices()]
    assert keys == ["general", "gov", "report", "thesis", "contract", "resume", "minutes"]


def test_table_numbering_skips_title_only_when_unique() -> None:
    """多章文档（有多个一级标题）里，第一个标题仍要参与编号。"""
    ir = parse_markdown("# 概述\n\n正文。\n\n# 方法\n\n正文。", title="文档")
    result = beautify_document(ir, template="report")
    headings = [block.text for block in result.ir.headings()]
    assert headings[0].startswith("1")
    assert headings[1].startswith("2")


# ---------------------------------------------------------------- 质量与体检


def test_analyze_document_reports_issues() -> None:
    ir = parse_markdown(
        "# 标题\n\n公司2024年度实现收入 1,200,000 元,同比增长 15%..  联系 13800138000 。\n\n"
        "重复段落内容足够长以便被识别为重复。\n\n重复段落内容足够长以便被识别为重复。\n",
        title="标题")
    analysis = analyze_document(ir)
    assert analysis["headings"] == 1
    assert analysis["duplicate_paragraphs"] == 1
    assert any("英文标点" in item for item in analysis["issues"])
    score = heuristic_score(ir)
    assert 0.0 <= score.overall <= 1.0
    assert score.weakest()[0]


def test_factual_consistency_penalises_number_changes() -> None:
    before = parse_plain_text("收入 100 元，成本 20 元。")
    after = parse_plain_text("收入 999 元，成本 20 元。")
    assert heuristic_score(after, before=before).factual_consistency < 1.0


# ---------------------------------------------------------------- 存储


def test_storage_documents_and_versions(tmp_path: Path) -> None:
    store = DocumentStorage(tmp_path / "modu.db")
    source = tmp_path / "a.md"
    source.write_text("# 标题\n\n正文。", encoding="utf-8")

    ir = parse_document(source)
    document_id = store.upsert_document(ir.path, title=ir.title, format_key=ir.format_key,
                                        blocks=len(ir.blocks))
    version_id = store.save_version(ir, document_id=document_id, label="打开", kind="open")
    assert version_id > 0

    record = store.get_document(document_id)
    assert record is not None and record.path == str(source)
    assert store.get_document_by_path(str(source)).id == document_id

    snapshot = store.load_version_ir(version_id)
    assert snapshot is not None and snapshot.title == ir.title
    assert store.toggle_favorite(document_id) is True
    assert store.list_documents(favorited_only=True)
    assert store.stats()["documents"] == 1

    store.add_audit("open", str(source), "测试")
    assert store.list_audit()[0]["action"] == "open"
    store.delete_document(document_id)
    assert store.list_versions(document_id) == []


def test_storage_settings_namespace(tmp_path: Path) -> None:
    store = DocumentStorage(tmp_path / "modu.db")
    store.set_setting("template", "gov")
    assert store.get_setting("template") == "gov"
    rows = store.query("SELECT key FROM app_settings")
    assert rows[0]["key"] == "document/template"


def test_save_remembers_new_path(tmp_path: Path) -> None:
    """另存为之后文档跟着新路径走（Ctrl+S 不再每次弹路径），并进文档库。"""
    store = DocumentStorage(tmp_path / "modu.db")
    library = DocumentLibrary(store, output_dir=tmp_path / "out")
    ir = DocumentIR.from_text("正文。", title="未命名")
    assert ir.path == ""

    target = tmp_path / "落盘.md"
    library.save(ir, target)
    assert ir.path == str(target)
    assert store.get_document_by_path(str(target)) is not None


def test_save_as_registers_new_document(tmp_path: Path) -> None:
    """另存到另一个文件：新文件单独登记，不与原文件共用文档记录。"""
    store = DocumentStorage(tmp_path / "modu.db")
    library = DocumentLibrary(store, output_dir=tmp_path / "out")
    source = tmp_path / "原稿.md"
    source.write_text("# 原稿\n\n正文。", encoding="utf-8")
    ir = library.open_path(source)

    copy = tmp_path / "副本.md"
    library.save(ir, copy)
    assert ir.path == str(copy)
    original = store.get_document_by_path(str(source))
    duplicated = store.get_document_by_path(str(copy))
    assert original is not None and duplicated is not None
    assert original.id != duplicated.id


def test_version_keep_follows_setting(tmp_path: Path) -> None:
    """「版本保留」设置必须真的生效（不能再写死常量 40）。"""
    from modu_workbench.core.document.storage import MAX_VERSIONS_PER_DOCUMENT

    store = DocumentStorage(tmp_path / "modu.db")
    library = DocumentLibrary(store, output_dir=tmp_path / "out")
    source = tmp_path / "a.md"
    source.write_text("# 标题\n\n正文。", encoding="utf-8")
    document = library.open_path(source)
    document_id = library.document_id_of(document)

    store.version_keep_provider = lambda: 3
    for index in range(8):
        document.blocks[-1].text = f"第 {index} 次修改"
        library.snapshot(document, label=f"第 {index} 次")
    assert store.version_keep() == 3
    assert len(store.list_versions(document_id, limit=50)) == 3

    # 读取器缺失或给出非法值时退回常量，不能把快照全删光
    store.version_keep_provider = None
    assert store.version_keep() == MAX_VERSIONS_PER_DOCUMENT
    store.version_keep_provider = lambda: 0
    assert store.version_keep() == MAX_VERSIONS_PER_DOCUMENT


def test_library_flow(tmp_path: Path) -> None:
    store = DocumentStorage(tmp_path / "modu.db")
    library = DocumentLibrary(store, output_dir=tmp_path / "out")
    source = tmp_path / "a.md"
    source.write_text("# 报告\n\n## 1.1 数据\n\n正文。", encoding="utf-8")

    ir = library.open_path(source)
    assert library.document_id_of(ir) > 0

    result = library.beautify(ir, template="report")
    exported = library.export(result.ir, "docx", name="美化结果")
    assert exported.path.is_file()

    saved = library.save(result.ir, source)
    assert saved.path.is_file()
    assert library.versions(ir)
    assert library.stats()["documents"] == 1

    snapshot_id = store.list_versions(library.document_id_of(ir))[0]["id"]
    assert library.rollback(snapshot_id) is not None
    assert "文档库" in library.describe()
