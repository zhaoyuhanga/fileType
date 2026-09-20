"""墨软文档：批注能力测试（需求 6.2「批注」）。

覆盖三层：
- 文档层：块的 `note` + 库的增删列（含审计）；
- 导出层：Word **原生批注**、Excel 批注工作表、HTML/Markdown/TXT 附录；
- PDF 层：真实 `/FreeText` 标注的写入与读取（pypdf）；
- 界面层：查看器批注面板与 PDF 工具箱「批注」页签。
"""
from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from modu_workbench.core.document import (
    TOOL_REGISTRY,
    BeautifyOptions,
    DocumentLibrary,
    DocumentStorage,
    annotate_pdf,
    call_tool,
    list_annotations,
    parse_document,
    parse_markdown,
    tool_names,
    write_document,
)

MARKDOWN = """# 季度报告

第一条结论需要补充数据。

第二条结论保持不动。

| 项目 | 金额 |
| --- | --- |
| A | 100 |
"""


@pytest.fixture()
def library(tmp_path: Path) -> DocumentLibrary:
    return DocumentLibrary(DocumentStorage(tmp_path / "modu.db"), output_dir=tmp_path / "out")


@pytest.fixture()
def ir():
    return parse_markdown(MARKDOWN, title="季度报告")


# ---------------------------------------------------------------- 文档层


def test_add_list_remove_comments(library: DocumentLibrary, ir) -> None:  # noqa: ANN001
    assert library.add_comment(ir, 1, "这里要补 2024 年数据", author="张三") is True
    comments = library.list_comments(ir)
    assert len(comments) == 1
    assert comments[0]["anchor"].startswith("第一条结论")
    assert comments[0]["note"] == "这里要补 2024 年数据"
    assert comments[0]["author"] == "张三"

    assert library.add_comment(ir, 2, "这条保留", author="李四") is True
    assert len(library.list_comments(ir)) == 2
    text = library.comments_text(ir)
    assert "批注汇总" in text and "张三" in text and "李四" in text

    assert library.remove_comment(ir, 1) is True
    assert len(library.list_comments(ir)) == 1
    assert library.clear_comments(ir) == 1
    assert library.list_comments(ir) == []

    actions = [record["action"] for record in library.audit_records()]
    assert "comment" in actions, "批注操作要写审计"


def test_comment_validation(library: DocumentLibrary, ir) -> None:  # noqa: ANN001
    assert library.add_comment(ir, 99, "越界") is False
    assert library.add_comment(ir, 1, "   ") is False
    assert library.remove_comment(ir, 1) is False, "没有批注时删除应返回 False"
    assert library.comments_text(ir) == ""


def test_comment_survives_snapshot_roundtrip(library: DocumentLibrary, ir) -> None:  # noqa: ANN001
    library.add_comment(ir, 1, "快照里也要有", author="张三")
    version_id = library.snapshot(ir, label="带批注")
    restored = library.rollback(version_id)
    assert restored is not None
    assert library.list_comments(restored)[0]["note"] == "快照里也要有"


# ---------------------------------------------------------------- 导出层


def test_word_export_writes_native_comments(library: DocumentLibrary, ir,  # noqa: ANN001
                                           tmp_path: Path) -> None:
    library.add_comment(ir, 1, "Word 原生批注验证", author="张三")
    result = write_document(ir, tmp_path / "带批注.docx", "docx",
                           options=BeautifyOptions(export_comments=True))
    with zipfile.ZipFile(result.path) as bundle:
        names = bundle.namelist()
        assert "word/comments.xml" in names, "应当写出 Word 原生批注部件"
        payload = bundle.read("word/comments.xml").decode("utf-8", "ignore")
    assert "Word 原生批注验证" in payload
    assert "张三" in payload


def test_word_export_without_comments_skips_part(ir, tmp_path: Path) -> None:  # noqa: ANN001
    result = write_document(ir, tmp_path / "无批注.docx", "docx")
    with zipfile.ZipFile(result.path) as bundle:
        assert "word/comments.xml" not in bundle.namelist()


def test_excel_export_writes_comment_sheet(library: DocumentLibrary, ir,  # noqa: ANN001
                                          tmp_path: Path) -> None:
    import openpyxl

    library.add_comment(ir, 1, "Excel 批注验证", author="李四")
    result = write_document(ir, tmp_path / "带批注.xlsx", "xlsx")
    assert any("批注" in note for note in result.notes)
    workbook = openpyxl.load_workbook(result.path)
    assert "批注" in workbook.sheetnames
    sheet = workbook["批注"]
    assert sheet.cell(row=1, column=1).value == "块序号"
    assert sheet.cell(row=2, column=3).value == "Excel 批注验证"
    workbook.close()


def test_text_formats_append_comment_appendix(library: DocumentLibrary, ir,  # noqa: ANN001
                                             tmp_path: Path) -> None:
    library.add_comment(ir, 1, "附录里的批注", author="张三")

    markdown = write_document(ir, tmp_path / "a.md", "markdown")
    markdown_text = markdown.path.read_text(encoding="utf-8")
    assert "## 批注" in markdown_text and "附录里的批注" in markdown_text

    html = write_document(ir, tmp_path / "a.html", "html")
    assert "附录里的批注" in html.path.read_text(encoding="utf-8")

    text = write_document(ir, tmp_path / "a.txt", "txt")
    assert "〔批注：附录里的批注〕" in text.path.read_text(encoding="utf-8")

    # 关掉开关后不再带批注
    plain = write_document(ir, tmp_path / "b.md", "markdown",
                           options=BeautifyOptions(export_comments=False))
    assert "附录里的批注" not in plain.path.read_text(encoding="utf-8")


# ---------------------------------------------------------------- PDF 批注


def _make_pdf(path: Path, pages: int = 2) -> Path:
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=595, height=842)
    with open(path, "wb") as handle:
        writer.write(handle)
    return path


def test_pdf_annotate_roundtrip(tmp_path: Path) -> None:
    source = _make_pdf(tmp_path / "空白.pdf", 2)
    produced = annotate_pdf(source, tmp_path / "批注.pdf",
                            [{"page": 1, "text": "第一条批注"},
                             (2, "第二条批注"),
                             {"page": 2, "text": "同页第二条"}])
    assert produced.is_file()

    found = list_annotations(produced)
    assert [item["text"] for item in found] == ["第一条批注", "第二条批注", "同页第二条"]
    assert [item["page"] for item in found] == [1, 2, 2]
    assert all(item["type"] == "FreeText" for item in found)


def test_pdf_annotate_errors(tmp_path: Path) -> None:
    from modu_workbench.core.document import PdfError

    source = _make_pdf(tmp_path / "a.pdf", 1)
    with pytest.raises(PdfError):
        annotate_pdf(source, tmp_path / "x.pdf", [])
    with pytest.raises(PdfError):
        annotate_pdf(source, tmp_path / "y.pdf", [{"page": 9, "text": "越界"}])
    assert list_annotations(source) == []


# ---------------------------------------------------------------- 工具与界面


def test_comment_tools(library: DocumentLibrary, ir) -> None:  # noqa: ANN001
    names = tool_names()
    assert {"add_comment", "list_comments", "annotate_pdf", "pdf_list_annotations"}.issubset(
        set(names))
    for name in ("add_comment", "list_comments"):
        assert TOOL_REGISTRY[name].handler is not None

    context = library.tool_context(ir)
    result = call_tool("add_comment", context, {"index": 1, "text": "工具加的批注"})
    assert result.ok and result.data["index"] == 1
    assert context.ir.blocks[1].note == "工具加的批注"

    # 用 anchor 定位
    result = call_tool("add_comment", context, {"anchor": "第二条结论", "text": "按文字定位"})
    assert result.ok and context.ir.blocks[2].note == "按文字定位"

    result = call_tool("add_comment", context, {"anchor": "不存在的文字", "text": "x"})
    assert result.ok is False and "没有找到" in result.summary

    result = call_tool("list_comments", context, {})
    assert result.ok and len(result.data["comments"]) == 2


def test_annotate_pdf_tool_uses_document_comments(library: DocumentLibrary,
                                                 tmp_path: Path) -> None:
    from modu_workbench.core.document import BLOCK_PARAGRAPH, Block

    pdf_path = _make_pdf(tmp_path / "扫描.pdf", 2)
    document = parse_document(pdf_path)
    # 空白 PDF 没有文本层：手工造一块带页码的正文，模拟"OCR 之后加批注"
    document.blocks = [Block(kind=BLOCK_PARAGRAPH, text="扫描页文字", meta={"page": 1})]
    assert library.add_comment(document, 0, "这份 PDF 的批注")
    context = library.tool_context(document)

    result = call_tool("annotate_pdf", context, {})
    assert result.ok and Path(result.data["path"]).is_file()
    assert list_annotations(result.data["path"])[0]["text"] == "这份 PDF 的批注"

    result = call_tool("pdf_list_annotations", context,
                       {"path": str(result.data["path"])})
    assert result.ok and "1 条批注" in result.summary


def test_viewer_comment_panel(qapp, tmp_path: Path) -> None:  # noqa: ANN001
    from modu_workbench.boards.document.viewer import DocViewer

    library = DocumentLibrary(DocumentStorage(tmp_path / "modu.db"), output_dir=tmp_path)
    document = parse_markdown(MARKDOWN, title="季度报告")
    viewer = DocViewer(library)
    viewer.load(document)
    try:
        item = viewer._outline.topLevelItem(0)          # noqa: SLF001
        viewer._outline.setCurrentItem(item)            # noqa: SLF001
        item.setSelected(True)
        assert library.add_comment(document, 0, "面板里的批注", author="张三")
        viewer._render_comments()                       # noqa: SLF001
        assert viewer._comment_list.count() == 1        # noqa: SLF001
        assert len(viewer.comments()) == 1

        viewer._jump_to_comment(viewer._comment_list.item(0))   # noqa: SLF001
        viewer._remove_comment()                        # noqa: SLF001
        qapp.processEvents()
        assert viewer._comment_list.count() == 0        # noqa: SLF001
    finally:
        # 清掉脏标记再关闭：否则 closeEvent 会弹"放弃未保存修改？"模态框（无头环境会一直等）
        viewer._set_dirty(False)                        # noqa: SLF001
        viewer._autosave_timer.stop()                   # noqa: SLF001
        viewer.close()


def test_pdf_dialog_annotation_tab(qapp, tmp_path: Path) -> None:  # noqa: ANN001
    from modu_workbench.boards.document.pdf_dialog import PdfToolsDialog

    source = _make_pdf(tmp_path / "空白.pdf", 2)
    library = DocumentLibrary(DocumentStorage(tmp_path / "modu.db"), output_dir=tmp_path / "out")
    dialog = PdfToolsDialog(library, source=str(source))
    try:
        labels = [dialog._tabs.tabText(index) for index in range(dialog._tabs.count())]  # noqa: SLF001
        assert "批注" in labels

        dialog._load_annotations()                      # noqa: SLF001
        assert "没有批注" in dialog._log.toPlainText()   # noqa: SLF001

        item = dialog._new_annotation_item()            # noqa: SLF001
        item.setText("2｜对话框写入的批注")
        dialog._annotations.addItem(item)               # noqa: SLF001
        dialog._do_annotate()                           # noqa: SLF001
        qapp.processEvents()
        assert "已写入 1 条 PDF 批注" in dialog._log.toPlainText()   # noqa: SLF001

        produced = sorted((tmp_path / "out").glob("*批注*.pdf"))
        assert produced, "应当在输出目录生成带批注的 PDF"
        assert list_annotations(produced[-1])[0]["text"] == "对话框写入的批注"

        # 再读一次应能读回
        dialog._source_field.setText(str(produced[-1]))  # noqa: SLF001
        dialog._load_annotations()                      # noqa: SLF001
        assert dialog._annotations.count() == 1         # noqa: SLF001
    finally:
        dialog.close()
