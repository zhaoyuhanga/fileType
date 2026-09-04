"""转换高级测试：Word(docx) / 表格(xlsx,xls,csv) / 媒体(ffmpeg 可选)。"""
from __future__ import annotations

from pathlib import Path

import pytest
from PySide6.QtWidgets import QApplication

from modu_workbench.core.convert.engine import run_conversion
from modu_workbench.core.convert.media_io import find_ffmpeg
from modu_workbench.core.convert.registry import get_action


@pytest.fixture(scope="module")
def qapp() -> QApplication:
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    return app


@pytest.fixture()
def out_dir(tmp_path: Path) -> Path:
    path = tmp_path / "out"
    path.mkdir()
    return path


def _make_docx(path: Path) -> None:
    import docx as docx_lib

    document = docx_lib.Document()
    document.add_heading("标题一", 1)
    document.add_paragraph("第一段正文内容。")
    document.add_paragraph("第二段内容。")
    document.save(path)


def test_docx_to_txt_html_pdf(qapp: QApplication, tmp_path: Path, out_dir: Path) -> None:
    docx = tmp_path / "示例.docx"
    _make_docx(docx)

    txt = run_conversion(get_action("docx-to-txt"), docx, out_dir)
    assert txt.status == "succeeded", txt.message
    content = (out_dir / "示例.txt").read_text(encoding="utf-8")
    assert "标题一" in content and "第一段正文内容" in content

    html = run_conversion(get_action("docx-to-html"), docx, out_dir)
    assert html.status == "succeeded"
    assert "<p>第一段正文内容。</p>" in (out_dir / "示例.html").read_text(encoding="utf-8")

    pdf = run_conversion(get_action("docx-to-pdf"), docx, out_dir)
    assert pdf.status == "succeeded", pdf.message
    assert (out_dir / "示例.pdf").read_bytes()[:4] == b"%PDF"


def test_doc_legacy_gives_guidance(tmp_path: Path, out_dir: Path) -> None:
    legacy = tmp_path / "旧版.doc"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0 legacy binary")  # OLE 头但内容不可解析
    result = run_conversion(get_action("doc-to-txt"), legacy, out_dir)
    assert result.status == "failed"
    assert "DOCX" in (result.message or "") or "LibreOffice" in (result.message or "")


def test_xlsx_to_csv_html_pdf(qapp: QApplication, tmp_path: Path, out_dir: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["名称", "数量"])
    ws.append(["苹果", 3])
    ws.append(["香蕉", 5])
    xlsx = tmp_path / "清单.xlsx"
    wb.save(xlsx)

    csv_res = run_conversion(get_action("xlsx-to-csv"), xlsx, out_dir)
    assert csv_res.status == "succeeded", csv_res.message
    csv_text = (out_dir / "清单.csv").read_text(encoding="utf-8")
    assert "名称,数量" in csv_text and "苹果,3" in csv_text

    html_res = run_conversion(get_action("xlsx-to-html"), xlsx, out_dir)
    assert html_res.status == "succeeded"
    assert "<table" in (out_dir / "清单.html").read_text(encoding="utf-8")

    pdf_res = run_conversion(get_action("xlsx-to-pdf"), xlsx, out_dir)
    assert pdf_res.status == "succeeded", pdf_res.message
    assert (out_dir / "清单.pdf").read_bytes()[:4] == b"%PDF"


def test_csv_to_html(tmp_path: Path, out_dir: Path) -> None:
    csv_file = tmp_path / "data.csv"
    csv_file.write_text("id,name\n1,墨读", encoding="utf-8")
    result = run_conversion(get_action("csv-to-html"), csv_file, out_dir)
    assert result.status == "succeeded", result.message
    assert "墨读" in (out_dir / "data.html").read_text(encoding="utf-8")


def test_media_conversion_when_ffmpeg_available(qapp: QApplication, tmp_path: Path, out_dir: Path) -> None:
    if find_ffmpeg() is None:
        pytest.skip("未安装 ffmpeg，跳过媒体转换用例")

    import subprocess

    wav = tmp_path / "tone.wav"
    ffmpeg = find_ffmpeg()
    subprocess.run(
        [ffmpeg, "-y", "-hide_banner", "-loglevel", "error",
         "-f", "lavfi", "-i", "sine=frequency=440:duration=0.2", str(wav)],
        check=True,
    )
    result = run_conversion(get_action("wav-to-mp3"), wav, out_dir)
    assert result.status == "succeeded", result.message
    assert (out_dir / "tone.mp3").stat().st_size > 0


def test_media_reports_missing_ffmpeg(tmp_path: Path, out_dir: Path) -> None:
    if find_ffmpeg() is not None:
        pytest.skip("已安装 ffmpeg，跳过缺失提示用例")
    fake = tmp_path / "x.wav"
    fake.write_bytes(b"RIFF fake")
    result = run_conversion(get_action("wav-to-mp3"), fake, out_dir)
    assert result.status == "failed"
    assert "ffmpeg" in (result.message or "")
