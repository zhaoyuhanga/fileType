"""表格读写：xlsx(openpyxl) / xls(xlrd) / ods(LibreOffice) / csv / tsv
→ csv / tsv / xlsx / markdown / html / txt / pdf。
"""
from __future__ import annotations

import csv
import html as html_mod
import io
from pathlib import Path

from .pdf_out import render_text_pdf

# 读表时按后缀分派
_DELIMITED = {".csv": ",", ".tsv": "\t"}


def _read_delimited(path: Path, delimiter: str) -> list[list[str]]:
    with open(path, encoding="utf-8-sig", newline="") as fp:
        return [[(cell or "") for cell in row] for row in csv.reader(fp, delimiter=delimiter)]


def _read_ods(path: Path) -> list[list[str]]:
    """ODS 没有随包依赖，交给 LibreOffice 转 CSV 再读；没有 soffice 时给可操作提示。"""
    from . import office_io

    generated = office_io.soffice_convert(path, target="csv")
    if generated is None:
        raise ValueError(
            "ODS 需要 LibreOffice（soffice）：请安装后重试，或先用 LibreOffice 另存为 XLSX / CSV"
        )
    try:
        return _read_delimited(generated, ",")
    finally:
        generated.unlink(missing_ok=True)


def read_sheet_rows(path: Path) -> list[list[str]]:
    """读取首个工作表为二维字符串列表。"""
    ext = path.suffix.lower()
    if ext in _DELIMITED:
        return _read_delimited(path, _DELIMITED[ext])
    if ext == ".ods":
        return _read_ods(path)
    if ext in (".xlsx", ".xlsm"):
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = wb.worksheets[0]
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(["" if cell is None else str(cell) for cell in row])
        wb.close()
        return rows
    if ext == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path))
        sheet = book.sheet_by_index(0)
        return [[sheet.cell_value(r, c) if sheet.cell_value(r, c) is not None else ""
                 for c in range(sheet.ncols)] for r in range(sheet.nrows)]
    raise ValueError("仅支持 XLSX / XLS / ODS / CSV / TSV 表格文件")


def _to_csv_text(rows: list[list[str]], delimiter: str = ",") -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=delimiter)
    writer.writerows(rows)
    return buffer.getvalue()


def _to_html(rows: list[list[str]], title: str) -> str:
    body = "\n".join(
        "<tr>" + "".join(f"<td>{html_mod.escape(c)}</td>" for c in row) + "</tr>" for row in rows
    )
    return (
        "<!doctype html><html lang='zh-CN'><head><meta charset='utf-8'>"
        f"<title>{html_mod.escape(title)}</title></head>"
        f"<body><table border='1' cellpadding='6' cellspacing='0'>{body}</table></body></html>"
    )


def _to_markdown(rows: list[list[str]]) -> str:
    """表格 → Markdown（首行当表头；列数不齐时补空单元格）。"""
    if not rows:
        return ""
    width = max(len(row) for row in rows)

    def line(cells: list[str]) -> str:
        padded = list(cells) + [""] * (width - len(cells))
        return "| " + " | ".join(cell.replace("|", "\\|") for cell in padded) + " |"

    header = line(rows[0])
    separator = "| " + " | ".join("---" for _ in range(width)) + " |"
    body = [line(row) for row in rows[1:]]
    return "\n".join([header, separator, *body]) + "\n"


def _to_xlsx(rows: list[list[str]], output: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    for row in rows:
        sheet.append(list(row))
    workbook.save(output)


def convert_sheet(path: Path, target: str, output: Path, title: str) -> None:
    rows = read_sheet_rows(path)
    if target == "csv":
        output.write_text(_to_csv_text(rows), encoding="utf-8")
        return
    if target == "tsv":
        output.write_text(_to_csv_text(rows, delimiter="\t"), encoding="utf-8")
        return
    if target == "xlsx":
        _to_xlsx(rows, output)
        return
    if target == "markdown":
        output.write_text(_to_markdown(rows), encoding="utf-8")
        return
    if target == "txt":
        text = "\n".join("\t".join(row) for row in rows)
        output.write_text(text, encoding="utf-8")
        return
    if target == "html":
        output.write_text(_to_html(rows, title), encoding="utf-8")
        return
    if target == "pdf":
        text = "\n".join("\t".join(row) for row in rows)
        render_text_pdf(text, output, title=title)
        return
    raise ValueError(f"不支持的表格目标：{target}")
