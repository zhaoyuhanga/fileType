"""墨软文档：文档解析层 —— 各格式 → 统一中间结构 `DocumentIR`。

分层约定：本模块**只做读取**（不写文件），任何格式解析失败都给中文可操作提示，
并尽量降级（例如 PDF 抽不到文字时提示 OCR），绝不让上层拿到半个对象。

格式分派（`parse_document`）：
- 纯文本族（txt/log/sql/代码）→ 段落 + 启发式标题；
- 标记族（md）→ 标题/列表/引用/代码块/表格/公式；
- 网页（html）→ bs4 遍历 h1-h6/p/li/blockquote/pre/table/img；
- 办公族（docx/odt/rtf/doc）→ python-docx 或内置 XML/RTF 解析，必要时 LibreOffice；
- 表格族（xlsx/xls/ods/csv/tsv）→ 二维表（保留公式原文）；
- 演示族（pptx/ppp/odp/dps）→ 逐页文本 + 备注；
- 数据族（json/xml/yaml/ini）→ 结构化美化文本；
- 电子书（epub）→ 章节 HTML 合并；
- 图片/扫描件 → 图片块 + 可选 OCR。

国产办公格式（wps/et/dps）先做 **OOXML 容器嗅探**，是 zip 就按对应族读，
否则退回 LibreOffice。
"""
from __future__ import annotations

import base64
import csv
import io
import json
import re
import zipfile
from pathlib import Path
from typing import Optional

from . import formats as fmt
from . import media
from .models import (
    BLOCK_CODE,
    BLOCK_HEADING,
    BLOCK_IMAGE,
    BLOCK_LIST,
    BLOCK_PAGE_BREAK,
    BLOCK_PARAGRAPH,
    BLOCK_QUOTE,
    BLOCK_SLIDE,
    BLOCK_TABLE,
    Block,
    DocumentIR,
    TableData,
)

#: 单个表格最多读入的行数（超出时截断并给出告警，避免大文件卡死界面）
MAX_TABLE_ROWS = 20000

#: 文本族扩展名 → 代码块语言（预览高亮用）
_CODE_LANGUAGES = {
    "py": "python", "pyw": "python", "js": "javascript", "mjs": "javascript",
    "cjs": "javascript", "ts": "typescript", "jsx": "jsx", "tsx": "tsx",
    "css": "css", "scss": "scss", "less": "less", "html": "html", "htm": "html",
    "xml": "xml", "json": "json", "sql": "sql", "java": "java", "c": "c",
    "h": "c", "cpp": "cpp", "hpp": "cpp", "go": "go", "rs": "rust",
    "php": "php", "rb": "ruby", "cs": "csharp", "kt": "kotlin", "swift": "swift",
    "sh": "bash", "bash": "bash", "zsh": "bash", "bat": "bat", "cmd": "bat",
    "ps1": "powershell", "vue": "html", "lua": "lua", "yaml": "yaml", "yml": "yaml",
    "ini": "ini", "toml": "toml", "log": "text", "md": "markdown",
}

#: 启发式标题：真正被 `guess_heading_level` 采纳的模式（顺序即优先级）
_HEADING_PATTERNS: tuple[tuple[re.Pattern, int], ...] = (
    (re.compile(r"^第[一二三四五六七八九十百千零〇\d]+[章篇部]"), 1),
    (re.compile(r"^[一二三四五六七八九十]+、"), 2),
    (re.compile(r"^[（(][一二三四五六七八九十]+[)）]"), 3),
    (re.compile(r"^\d+\.\d+\.\d+[\.、]?\s"), 4),
    (re.compile(r"^\d+\.\d+[\.、]?\s"), 3),
    (re.compile(r"^\d+[\.、]\s"), 2),
)

#: 形似「第一章 概述」但不希望被当成正文句子的行尾标点
_SENTENCE_END = ("。", "，", "；", ".", ",", ";")


class ParseError(ValueError):
    """解析失败（消息为可直接展示给用户的中文提示）。"""


# ---------------------------------------------------------------- 入口


def parse_document(path: str | Path, *, ocr: bool = False, max_rows: int = MAX_TABLE_ROWS) -> DocumentIR:
    """把任意受支持的文件解析成 `DocumentIR`。"""
    target = Path(path)
    if not target.is_file():
        raise ParseError(f"文件不存在：{target}")
    spec = fmt.spec_for_path(target)
    if spec is None:
        raise ParseError(
            f"暂不支持的文件类型：{target.suffix or target.name}。"
            "可在「墨软转换」里先转成受支持格式，或把它加进 formats.py 的格式表。"
        )

    if spec.key in ("docx", "doc", "wps"):
        ir = _parse_word_family(target, spec.key)
    elif spec.key in ("rtf", "odt"):
        ir = _parse_word_family(target, spec.key)
    elif spec.key == "pdf":
        ir = _parse_pdf(target, ocr=ocr)
    elif spec.key == "markdown":
        ir = parse_markdown(_read_text(target), title=target.stem, base_dir=target.parent)
    elif spec.key == "html":
        ir = parse_html(_read_text(target), title=target.stem, base_dir=target.parent)
    elif spec.key == "epub":
        ir = _parse_epub(target)
    elif spec.key in ("xlsx", "xls", "ods", "et"):
        ir = _parse_sheet_family(target, spec.key, max_rows=max_rows)
    elif spec.key in ("csv", "tsv"):
        ir = _parse_delimited(target, spec.key, max_rows=max_rows)
    elif spec.key in ("pptx", "ppt", "odp", "dps"):
        ir = _parse_slides_family(target, spec.key)
    elif spec.key in ("json", "xml", "yaml", "ini"):
        ir = _parse_data_family(target, spec.key)
    elif spec.key == "image" or spec.render == fmt.RENDER_IMAGE:
        ir = _parse_image(target, spec.key, ocr=ocr)
    else:
        ir = _parse_plain_text(target, spec.key)

    ir.path = str(target)
    ir.format_key = spec.key
    if not ir.title:
        ir.title = target.stem
    # 标题是"文件名"还是"文档自己的"要分清楚：导出时不给文件名造一行可见标题
    # （否则"打开 Word → 保存"会平白多出一行标题，见 writer._title_already_in_blocks）
    ir.metadata.setdefault(
        "title_source", "filename" if ir.title == target.stem else "content")
    try:
        stat = target.stat()
        ir.metadata.setdefault("size_bytes", stat.st_size)
        ir.metadata.setdefault("modified_at", int(stat.st_mtime))
    except OSError:
        pass
    ir.metadata.setdefault("format", spec.key)
    ir.metadata.setdefault("format_label", spec.label)
    ir.metadata.setdefault("category", spec.category_label)
    return ir


def read_document_text(path: str | Path) -> str:
    """只看文本的快捷入口（AI 上下文 / 摘要用）。"""
    return parse_document(path).text()


# ---------------------------------------------------------------- 工具


def _read_text(path: Path) -> str:
    from modu_workbench.core.convert import text_io

    try:
        return text_io.read_text_smart(path)
    except Exception as error:  # noqa: BLE001
        raise ParseError(f"读取文本失败：{error}") from error


def _looks_like_zip(path: Path) -> bool:
    """只读头 2 字节判断 zip 容器（不把整个文件读进内存）。"""
    try:
        with open(path, "rb") as handle:
            return handle.read(2) == b"PK"
    except OSError:
        return False


def sniff_ooxml(path: Path) -> str:
    """zip 容器嗅探：返回 docx / xlsx / pptx，非 OOXML 返回空串。"""
    try:
        with zipfile.ZipFile(path) as bundle:
            names = set(bundle.namelist())
    except (zipfile.BadZipFile, OSError):
        return ""
    if "word/document.xml" in names:
        return "docx"
    if any(name.startswith("xl/worksheets/") for name in names):
        return "xlsx"
    if any(name.startswith("ppt/slides/slide") for name in names):
        return "pptx"
    return ""


def guess_heading_level(text: str, *, numbered: bool = True) -> int:
    """启发式判断一行是不是标题（返回 0 表示不是）。

    规则保守：只认显式编号（第X章 / 一、 / 1.1）——
    "看起来短" 不作为依据，否则整篇正文会被误判成标题。
    """
    stripped = (text or "").strip()
    if not stripped or len(stripped) > 60:
        return 0
    if not numbered:
        return 0
    for pattern, level in _HEADING_PATTERNS:
        if pattern.match(stripped):
            # 编号后面跟着句末标点说明这是正文句子而非标题
            if stripped.endswith(_SENTENCE_END):
                return 0
            return level
    return 0


def join_wrapped_lines(lines: list[str]) -> str:
    """把 Markdown/文本里被硬换行拆开的行合并成一段（中文不留空格）。"""
    text = " ".join(line.strip() for line in lines if line.strip())
    return re.sub(r"(?<=[\u3400-\u9fff\u3000-\u303f])\s+(?=[\u3400-\u9fff\u3000-\u303f])", "", text)


# ---------------------------------------------------------------- 纯文本


def parse_plain_text(text: str, *, title: str = "", infer_headings: bool = True) -> DocumentIR:
    """纯文本 → 段落（空行分块）；命中编号模式的行升级成标题。"""
    blocks: list[Block] = []
    for chunk in re.split(r"\n\s*\n", text.replace("\r\n", "\n").replace("\r", "\n")):
        chunk = chunk.strip("\n")
        if not chunk.strip():
            continue
        lines = chunk.split("\n")
        if len(lines) == 1:
            line = lines[0].strip()
            level = guess_heading_level(line) if infer_headings else 0
            blocks.append(Block(kind=BLOCK_HEADING, text=line, level=level) if level
                          else Block(kind=BLOCK_PARAGRAPH, text=line))
            continue
        current: list[str] = []
        for line in lines:
            level = guess_heading_level(line) if infer_headings else 0
            if level and not current:
                blocks.append(Block(kind=BLOCK_HEADING, text=line.strip(), level=level))
                continue
            if level and current:
                blocks.append(Block(kind=BLOCK_PARAGRAPH, text=join_wrapped_lines(current)))
                current = []
                blocks.append(Block(kind=BLOCK_HEADING, text=line.strip(), level=level))
                continue
            if not line.strip():
                if current:
                    blocks.append(Block(kind=BLOCK_PARAGRAPH, text=join_wrapped_lines(current)))
                    current = []
                continue
            current.append(line)
        if current:
            blocks.append(Block(kind=BLOCK_PARAGRAPH, text=join_wrapped_lines(current)))
    return DocumentIR(blocks=blocks, title=title, format_key="txt")


def _parse_plain_text(path: Path, format_key: str) -> DocumentIR:
    text = _read_text(path)
    ir = parse_plain_text(text, title=path.stem)
    ir.format_key = format_key
    language = _CODE_LANGUAGES.get(path.suffix.lower().lstrip("."), "")
    if format_key in ("sql", "code", "log") and language:
        ir.metadata["language"] = language
        ir.styles["language"] = language
    if format_key == "log":
        ir.metadata["lines"] = len(text.splitlines())
    return ir


# ---------------------------------------------------------------- Markdown

_MD_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)+\|?\s*$")
_MD_LIST = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s+(.*)$")
_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_MD_QUOTE = re.compile(r"^\s*>\s?(.*)$")
_MD_IMAGE = re.compile(r"^\s*!\[([^\]]*)\]\(\s*<?([^)\s>]+)>?(?:\s+[\"'][^\"']*[\"'])?\s*\)\s*$")

#: `data:image/png;base64,....`
_DATA_URI = re.compile(r"^data:(image/[a-zA-Z0-9.+-]+)?;base64,(.*)$", re.DOTALL)


def _decode_data_uri(src: str) -> Optional[bytes]:
    """把 `data:image/png;base64,…` 解成字节（不是 data URI 或解不开时返回 None）。"""
    match = _DATA_URI.match((src or "").strip())
    if match is None:
        return None
    try:
        return base64.b64decode(match.group(2), validate=False)
    except Exception:  # noqa: BLE001  坏数据当"取不到"，交给调用方留占位
        return None


def _local_image(src: str, base_dir: Optional[Path]) -> Optional[Path]:
    """相对/绝对路径的本地图片（远程 URL 与 data URI 不算）。"""
    text = (src or "").strip().split("?")[0].split("#")[0]
    if not text or text.startswith(("http://", "https://", "//", "data:")):
        return None
    candidate = Path(text)
    if not candidate.is_absolute() and base_dir is not None:
        candidate = base_dir / candidate
    try:
        return candidate if candidate.is_file() else None
    except OSError:
        return None


def _image_block(alt: str, src: str, *, base_dir: Optional[Path] = None,
                 label: str = "") -> Block:
    """图片引用 → image 块。

    - data URI 与本地相对路径：字节收进素材缓存（`media.py`），块上只留引用；
    - 远程 URL：原样记在 `meta["src"]`（导出 HTML/Markdown 时照写，不联网抓图）。
    """
    text = (alt or "").strip() or Path((src or "").split("?")[0]).name or label or "图片"
    inline = _decode_data_uri(src)
    if inline:
        try:
            info = media.store_image(inline, name=text if Path(text).suffix else "image.png")
            return Block(kind=BLOCK_IMAGE, text=info["name"], meta={"image": info})
        except media.MediaError:
            return Block(kind=BLOCK_IMAGE, text=text, meta={"src": "data:"})
    local = _local_image(src, base_dir)
    if local is not None:
        try:
            info = media.store_image(local.read_bytes(), name=local.name)
            return Block(kind=BLOCK_IMAGE, text=info["name"], meta={"image": info})
        except (media.MediaError, OSError):
            return Block(kind=BLOCK_IMAGE, text=text, meta={"src": str(local)})
    return Block(kind=BLOCK_IMAGE, text=text, meta={"src": src})


def parse_markdown(text: str, *, title: str = "", base_dir: Optional[Path] = None) -> DocumentIR:
    """Markdown → IR（标题/列表/引用/代码块/表格/图片/公式）。

    `base_dir` 是源文件所在目录：`![](图.png)` 这类相对路径按它找到文件并收进素材缓存。
    """
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[Block] = []
    index = 0
    pending: list[str] = []

    def flush() -> None:
        if pending:
            joined = join_wrapped_lines(pending)
            if joined.strip():
                blocks.append(Block(kind=BLOCK_PARAGRAPH, text=joined.strip()))
            pending.clear()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```") or stripped.startswith("~~~"):
            flush()
            fence = stripped[:3]
            language = stripped[3:].strip()
            index += 1
            body: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith(fence):
                body.append(lines[index])
                index += 1
            index += 1
            blocks.append(Block(kind=BLOCK_CODE, text="\n".join(body).rstrip(),
                                style={"language": language} if language else {}))
            continue

        # 表格
        if "|" in line and index + 1 < len(lines) and _MD_TABLE_SEP.match(lines[index + 1]):
            flush()
            rows: list[list[str]] = [_split_md_row(line)]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append(_split_md_row(lines[index]))
                index += 1
            blocks.append(Block(kind=BLOCK_TABLE, table=TableData(rows=rows, header=True)))
            continue

        # 标题
        match = _MD_HEADING.match(line)
        if match:
            flush()
            blocks.append(Block(kind=BLOCK_HEADING, text=match.group(2).strip(),
                                level=len(match.group(1))))
            index += 1
            continue

        # 分隔线 → 分页
        if stripped in ("---", "***", "___"):
            flush()
            blocks.append(Block(kind=BLOCK_PAGE_BREAK))
            index += 1
            continue

        # 图片：![alt](src)（本地相对路径与 data URI 都收进素材缓存，远程 URL 原样保留）
        image = _MD_IMAGE.match(line)
        if image:
            flush()
            blocks.append(_image_block(image.group(1), image.group(2), base_dir=base_dir))
            index += 1
            continue

        # 引用
        quote = _MD_QUOTE.match(line)
        if quote:
            flush()
            body = [quote.group(1)]
            index += 1
            while index < len(lines) and _MD_QUOTE.match(lines[index]):
                body.append(_MD_QUOTE.match(lines[index]).group(1))     # type: ignore[union-attr]
                index += 1
            blocks.append(Block(kind=BLOCK_QUOTE, text=join_wrapped_lines(body)))
            continue

        # 列表
        item = _MD_LIST.match(line)
        if item:
            flush()
            while index < len(lines):
                hit = _MD_LIST.match(lines[index])
                if not hit:
                    break
                blocks.append(Block(kind=BLOCK_LIST, text=hit.group(1).strip()))
                index += 1
            continue

        if not stripped:
            flush()
            index += 1
            continue

        pending.append(line)
        index += 1

    flush()
    return DocumentIR(blocks=blocks, title=title, format_key="markdown")


def _split_md_row(line: str) -> list[str]:
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]
    return [cell.strip().replace("\\|", "|") for cell in text.split("|")]


# ---------------------------------------------------------------- HTML


def parse_html(text: str, *, title: str = "", base_dir: Optional[Path] = None) -> DocumentIR:
    """HTML → IR（h1-h6 / p / li / blockquote / pre / table / img）。

    `base_dir` 是源文件所在目录：相对路径的 `<img src>` 按它找到文件并收进素材缓存。
    """
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(text, "html.parser")
    for tag in soup(["script", "style", "noscript", "meta", "link"]):
        tag.decompose()
    if not title and soup.title and soup.title.string:
        title = soup.title.string.strip()

    body = soup.body or soup
    blocks: list[Block] = []

    def walk(node) -> None:  # noqa: ANN001
        name = getattr(node, "name", None)
        if name in (f"h{level}" for level in range(1, 7)):
            level = int(name[1])
            content = _clean_text(node.get_text(" ", strip=True))
            if content:
                blocks.append(Block(kind=BLOCK_HEADING, text=content, level=level))
            return
        if name == "p":
            content = _clean_text(node.get_text(" ", strip=True))
            if content:
                blocks.append(Block(kind=BLOCK_PARAGRAPH, text=content))
            return
        if name == "li":
            content = _clean_text(node.get_text(" ", strip=True))
            if content:
                blocks.append(Block(kind=BLOCK_LIST, text=content))
            return
        if name == "blockquote":
            content = _clean_text(node.get_text(" ", strip=True))
            if content:
                blocks.append(Block(kind=BLOCK_QUOTE, text=content))
            return
        if name in ("pre", "code") and getattr(node, "parent", None) is not None:
            if name == "pre" or node.parent.name not in ("pre",):
                blocks.append(Block(kind=BLOCK_CODE, text=node.get_text("\n").rstrip()))
            return
        if name == "table":
            rows = _html_table_rows(node)
            if rows:
                blocks.append(Block(kind=BLOCK_TABLE, table=TableData(rows=rows, header=True)))
            return
        if name == "img":
            blocks.append(_image_block(node.get("alt") or "", node.get("src") or "",
                                       base_dir=base_dir, label="图片"))
            return
        if name == "hr":
            blocks.append(Block(kind=BLOCK_PAGE_BREAK))
            return
        for child in getattr(node, "children", []):
            if getattr(child, "name", None):
                walk(child)

    walk(body)
    return DocumentIR(blocks=blocks, title=title, format_key="html")


def _html_table_rows(table) -> list[list[str]]:  # noqa: ANN001
    rows: list[list[str]] = []
    for tr in table.find_all("tr"):
        cells = tr.find_all(["td", "th"])
        if not cells:
            continue
        rows.append([_clean_text(cell.get_text(" ", strip=True)) for cell in cells])
    return rows


def _clean_text(text: str) -> str:
    return re.sub(r"[ \t\u00a0]+", " ", (text or "").replace("\n", " ")).strip()


# ---------------------------------------------------------------- 表格族


def _parse_delimited(path: Path, format_key: str, *, max_rows: int) -> DocumentIR:
    delimiter = "\t" if format_key == "tsv" else ","
    text = _read_text(path)
    rows: list[list[str]] = []
    truncated = 0
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    for index, row in enumerate(reader):
        if index >= max_rows:
            truncated += 1
            continue
        rows.append([cell for cell in row])
    if not rows:
        raise ParseError(f"表格是空的：{path.name}")
    table = TableData(rows=rows, name=path.stem, header=True,
                      meta={"delimiter": delimiter, "truncated_rows": truncated})
    ir = DocumentIR(blocks=[Block(kind=BLOCK_TABLE, table=table)], title=path.stem,
                    format_key=format_key)
    if truncated:
        ir.warnings.append(f"表格超过 {max_rows} 行，已截断（仅影响预览，导出仍按当前内容）")
    return ir


def _parse_sheet_family(path: Path, format_key: str, *, max_rows: int) -> DocumentIR:
    if format_key in ("xlsx", "et"):
        container = sniff_ooxml(path)
        if container == "xlsx":
            return _parse_xlsx(path, max_rows=max_rows)
        if container == "":
            pass        # 不是 zip：可能是旧 WPS 表格 → 交给 LibreOffice
    if format_key == "xls":
        return _parse_xls(path, max_rows=max_rows)
    converted = _soffice_convert(path, "xlsx")
    if converted is not None:
        try:
            return _parse_xlsx(converted, max_rows=max_rows)
        finally:
            converted.unlink(missing_ok=True)
    converted = _soffice_convert(path, "csv")
    if converted is not None:
        try:
            ir = _parse_delimited(converted, "csv", max_rows=max_rows)
            ir.title = path.stem
            ir.format_key = format_key
            return ir
        finally:
            converted.unlink(missing_ok=True)
    raise ParseError(
        f"{fmt.format_label(format_key)} 需要 LibreOffice（soffice）：请安装后重试，"
        "或先用 LibreOffice 另存为 XLSX / CSV。"
    )


def _parse_xlsx(path: Path, *, max_rows: int) -> DocumentIR:
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    except Exception as error:  # noqa: BLE001
        raise ParseError(f"无法读取 Excel：{error}") from error

    blocks: list[Block] = []
    warnings: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            truncated = 0
            formulas = 0
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= max_rows:
                    truncated += 1
                    continue
                cells = []
                for cell in row:
                    if cell is None:
                        cells.append("")
                        continue
                    text = str(cell)
                    if text.startswith("="):
                        formulas += 1
                    cells.append(text)
                rows.append(cells)
            if not rows:
                continue
            if truncated:
                warnings.append(f"工作表「{sheet.title}」超过 {max_rows} 行，已截断预览")
            table = TableData(
                rows=rows, name=sheet.title, header=True,
                meta={
                    "truncated_rows": truncated,
                    "formula_cells": formulas,
                    "freeze_panes": str(getattr(sheet, "freeze_panes", "") or ""),
                    "data_validations": len(getattr(sheet, "data_validations", []) or []),
                    "conditional_formats": sum(
                        len(getattr(item, "rules", []) or [])
                        for item in (getattr(sheet, "conditional_formatting", []) or [])
                    ),
                    "merged_cells": len(getattr(getattr(sheet, "merged_cells", None), "ranges", []) or []),
                },
            )
            blocks.append(Block(kind=BLOCK_TABLE, table=table))
    finally:
        try:
            workbook.close()
        except Exception:  # noqa: BLE001
            pass

    if not blocks:
        raise ParseError(f"表格里没有可读数据：{path.name}")
    ir = DocumentIR(blocks=blocks, title=path.stem, format_key="xlsx", warnings=warnings)
    return ir


def _parse_xls(path: Path, *, max_rows: int) -> DocumentIR:
    import xlrd

    try:
        book = xlrd.open_workbook(str(path))
    except Exception as error:  # noqa: BLE001
        raise ParseError(f"无法读取旧版 Excel：{error}") from error
    blocks: list[Block] = []
    for sheet in book.sheets():
        rows: list[list[str]] = []
        for row_index in range(min(sheet.nrows, max_rows)):
            rows.append([
                _cell_text(sheet.cell_value(row_index, column))
                for column in range(sheet.ncols)
            ])
        if rows:
            blocks.append(Block(kind=BLOCK_TABLE,
                                table=TableData(rows=rows, name=sheet.name, header=True)))
    if not blocks:
        raise ParseError(f"表格里没有可读数据：{path.name}")
    return DocumentIR(blocks=blocks, title=path.stem, format_key="xls")


def _cell_text(value) -> str:  # noqa: ANN001
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def table_from_rows(rows: list[list[str]], *, name: str = "", header: bool = True) -> TableData:
    return TableData(rows=[list(row) for row in rows], name=name, header=header)


# ---------------------------------------------------------------- Word 族


def _parse_word_family(path: Path, format_key: str) -> DocumentIR:
    suffix = path.suffix.lower()
    if format_key == "wps" or suffix in (".wps", ".docm"):
        container = sniff_ooxml(path)
        if container == "docx":
            return _parse_docx(path)
        if container == "xlsx":
            return _parse_xlsx(path, max_rows=MAX_TABLE_ROWS)
        if container == "pptx":
            return _parse_pptx(path)
    if suffix == ".docx" or (suffix == ".doc" and _looks_like_zip(path)):
        return _parse_docx(path)

    if suffix == ".odt" or format_key == "odt":
        try:
            return _parse_odt(path)
        except ParseError:
            pass        # 交给 LibreOffice
    if suffix == ".rtf" or format_key == "rtf":
        try:
            return _parse_rtf(path)
        except ParseError:
            pass

    converted = _soffice_convert(path, "docx")
    if converted is not None:
        try:
            ir = _parse_docx(converted)
            ir.title = path.stem
            ir.format_key = format_key
            return ir
        finally:
            converted.unlink(missing_ok=True)
    converted = _soffice_convert(path, "txt")
    if converted is not None:
        try:
            ir = parse_plain_text(_read_text(converted), title=path.stem)
            ir.format_key = format_key
            ir.warnings.append("通过 LibreOffice 抽取纯文本：复杂排版（表格/图片）未保留")
            return ir
        finally:
            converted.unlink(missing_ok=True)
    raise ParseError(
        f"{fmt.format_label(format_key)} 需要 LibreOffice（soffice）才能读取："
        "请安装后重试，或先另存为 DOCX。"
    )


def _parse_docx(path: Path) -> DocumentIR:
    import docx as docx_lib
    from docx.oxml.ns import qn
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph

    try:
        document = docx_lib.Document(str(path))
    except Exception as error:  # noqa: BLE001
        raise ParseError(f"无法读取 Word 文档：{error}") from error

    blocks: list[Block] = []
    warnings: list[str] = []
    body = document.element.body
    if body.find(".//" + qn("w:ins")) is not None or body.find(".//" + qn("w:del")) is not None:
        warnings.append("文档含修订标记：按「接受全部修订」后的结果读取，被删除的内容未保留")
    textbox_seen = False
    for child in body.iterchildren():
        tag = child.tag
        if tag == qn("w:p"):
            paragraph = DocxParagraph(child, document)
            # 不用 `paragraph.text`：它只看直接子级 w:r，
            # 修订插入（w:ins）与超链接（w:hyperlink）里的文字会被漏掉
            text = _docx_paragraph_text(child).strip()
            style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
            level = _heading_level_from_style(style_name)
            if text:
                block = Block(
                    kind=BLOCK_HEADING if level else BLOCK_PARAGRAPH,
                    text=text,
                    level=level or 0,
                    style={"style_name": style_name} if style_name else {},
                )
                if level:
                    block.style["style_name"] = style_name
                blocks.append(block)
            for inner in _docx_textbox_paragraphs(child):
                box_text = _docx_paragraph_text(inner).strip()
                if box_text:
                    # 文本框在 IR 里没有"浮动位置"，退化为引用块而不是丢掉
                    blocks.append(Block(kind=BLOCK_QUOTE, text=box_text,
                                        style={"textbox": True}))
                    textbox_seen = True
            for image in _docx_images(paragraph, document, warnings):
                # 图片不进 IR 正文，只留引用（见 `media.py`）：写出时重新嵌入
                blocks.append(Block(kind=BLOCK_IMAGE, text=image["name"],
                                    meta={"image": image}))
        elif tag == qn("w:tbl"):
            table = DocxTable(child, document)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if rows:
                blocks.append(Block(kind=BLOCK_TABLE, table=TableData(
                    rows=rows, header=True,
                    styles={"style_name": (table.style.name if table.style is not None else "") or ""},
                )))
            # 表格单元格里的图片 IR 装不下位置，退化为"表格后面跟一张图"而不是丢掉
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for image in _docx_images(paragraph, document, warnings):
                            blocks.append(Block(kind=BLOCK_IMAGE,
                                                text=f"{image['name']}（表格内图片）",
                                                meta={"image": image, "from_table": True}))

    metadata: dict = {}
    try:
        props = document.core_properties
        metadata = {
            "author": props.author or "",
            "last_modified_by": props.last_modified_by or "",
            "created": str(props.created or ""),
            "modified": str(props.modified or ""),
            "revision": int(props.revision or 0),
            "subject": props.subject or "",
        }
    except Exception:  # noqa: BLE001
        metadata = {}
    for key, text in _docx_header_footer(document).items():
        if text:
            metadata[key] = text
    if textbox_seen:
        warnings.append("文本框里的文字已按引用块保留（Word 里是浮动文本框，位置不再保留）")
    title = str(metadata.get("title") or "") or path.stem
    ir = DocumentIR(blocks=blocks, title=title, format_key="docx", metadata=metadata)
    ir.warnings.extend(warnings)
    return ir


def _docx_paragraph_text(element: object) -> str:
    """段落文字：按文档顺序取全部 `w:t`，跳过被修订删除的内容与文本框正文。

    为什么不用 `paragraph.text`：它只拼接**直接子级** `w:r`，于是
    `w:ins`（修订插入）、`w:hyperlink`（超链接）、`w:smartTag` 里的文字会整段消失。
    """
    from docx.oxml.ns import qn

    skipped = (qn("w:del"), qn("w:txbxContent"))
    parts: list[str] = []
    for node in element.iter():                    # type: ignore[attr-defined]
        tag = node.tag
        if tag == qn("w:t"):
            if _docx_inside(node, element, skipped):
                continue
            parts.append(node.text or "")
        elif tag == qn("w:tab"):
            parts.append("\t")
        elif tag in (qn("w:br"), qn("w:cr")):
            parts.append("\n")
    return "".join(parts)


def _docx_inside(node: object, top: object, tags: tuple) -> bool:
    """`node` 是否位于 `top` 之内、且中途经过 `tags` 里的容器。"""
    current = node.getparent()                     # type: ignore[attr-defined]
    while current is not None and current is not top:
        if current.tag in tags:
            return True
        current = current.getparent()
    return False


def _docx_textbox_paragraphs(element: object) -> list[object]:
    """段落里的文本框内容（`w:txbxContent` 下的段落）。"""
    from docx.oxml.ns import qn

    boxes = []
    for content in element.iter(qn("w:txbxContent")):   # type: ignore[attr-defined]
        boxes.extend(content.iter(qn("w:p")))
    return boxes


def _docx_header_footer(document: object) -> dict:
    """页眉页脚文字（正文之外的内容，收进 metadata 而不是当正文块）。"""
    result: dict = {}
    try:
        sections = list(getattr(document, "sections", []) or [])
    except Exception:  # noqa: BLE001
        return result
    for part_name, attr in (("header_text", "header"), ("footer_text", "footer")):
        lines: list[str] = []
        for section in sections:
            try:
                holder = getattr(section, attr)
                lines.extend(paragraph.text.strip() for paragraph in holder.paragraphs
                             if paragraph.text.strip())
            except Exception:  # noqa: BLE001  链接到上一节/没有页眉部件时跳过
                continue
        if lines:
            result[part_name] = " / ".join(dict.fromkeys(lines))     # 去重但保持顺序
    return result


def _docx_images(paragraph: object, document: object, warnings: list[str]) -> list[dict]:
    """段落里的图片：DrawingML（`w:drawing`）与老式 VML（`v:imagedata`）都认。

    返回 `media.store_image()` 的引用（含文件名/尺寸）；存不下时往 `warnings` 里
    记一条原因，而不是静默丢图。
    """
    from docx.oxml.ns import qn

    element = getattr(paragraph, "_element", None)
    if element is None:
        return []
    images: list[dict] = []
    for drawing in element.iter(qn("w:drawing")):
        width_pt = height_pt = 0.0
        extent = drawing.find(".//" + qn("wp:extent"))
        if extent is not None:
            try:
                width_pt = int(extent.get("cx") or 0) / 12700.0     # EMU → pt
                height_pt = int(extent.get("cy") or 0) / 12700.0
            except (TypeError, ValueError):
                width_pt = height_pt = 0.0
        hint = _docx_image_hint(drawing)
        for blip in drawing.iter(qn("a:blip")):
            image = _docx_image_ref(document, blip.get(qn("r:embed")) or blip.get(qn("r:link")),
                                    warnings, width_pt=width_pt, height_pt=height_pt, name_hint=hint)
            if image:
                images.append(image)
    for shape in element.iter(_VML_SHAPE):
        # 老式 VML 图片（Word 97-2003 兼容写法）：尺寸写在 style 里
        style = str(shape.get("style") or "")
        width_pt = _vml_length(style, "width")
        height_pt = _vml_length(style, "height")
        for imagedata in shape.iter(_VML_IMAGEDATA):
            image = _docx_image_ref(document, imagedata.get(qn("r:id")), warnings,
                                    width_pt=width_pt, height_pt=height_pt)
            if image:
                images.append(image)
    return images


#: VML 命名空间不在 python-docx 的 nsmap 里，直接用 Clark 记法
_VML_SHAPE = "{urn:schemas-microsoft-com:vml}shape"
_VML_IMAGEDATA = "{urn:schemas-microsoft-com:vml}imagedata"


def _vml_length(style: str, name: str) -> float:
    """从 VML 的 `style="width:120pt;height:90pt"` 里取 pt 尺寸（认不出返回 0）。"""
    match = re.search(rf"{name}\s*:\s*([0-9.]+)\s*pt", style or "", re.IGNORECASE)
    return float(match.group(1)) if match else 0.0


def _docx_image_hint(drawing: object) -> str:
    """图片的"人话名"：`wp:docPr@descr` 里带扩展名时用它（Word 会写原始文件名）。

    取不到就返回空串，让调用方退回部件名（`image1.png` 这种）。
    """
    from docx.oxml.ns import qn

    doc_pr = drawing.find(".//" + qn("wp:docPr"))      # type: ignore[attr-defined]
    if doc_pr is None:
        return ""
    descr = str(doc_pr.get("descr") or "").strip()
    if descr and Path(descr).suffix:
        return Path(descr).name
    return ""


def _docx_image_ref(document: object, rel_id: Optional[str], warnings: list[str], *,
                    width_pt: float = 0.0, height_pt: float = 0.0,
                    name_hint: str = "") -> Optional[dict]:
    """按关系 id 取出图片部件并落进素材缓存。"""
    if not rel_id:
        return None
    part = getattr(document, "part", None)
    try:
        related = part.related_parts[rel_id]      # type: ignore[index]
    except (KeyError, AttributeError, TypeError):
        return None
    blob = getattr(related, "blob", b"") or b""
    if not blob:
        return None
    part_name = str(getattr(related, "partname", "") or "")
    try:
        return media.store_image(blob, content_type=str(getattr(related, "content_type", "") or ""),
                                 name=name_hint or Path(part_name).name or "image",
                                 width_pt=width_pt, height_pt=height_pt)
    except media.MediaError as error:
        warnings.append(f"Word 里有张图片没能保留：{error}")
        return None


def _heading_level_from_style(style_name: str) -> int:
    if not style_name:
        return 0
    lowered = style_name.strip().lower()
    match = re.search(r"(?:heading|标题)\s*(\d)", lowered)
    if match:
        return max(1, min(6, int(match.group(1))))
    if lowered in ("title", "标题"):
        return 1
    if lowered in ("subtitle", "副标题"):
        return 2
    return 0


def _parse_odt(path: Path) -> DocumentIR:
    """ODT：纯标准库读 content.xml（段落/标题/表格）。"""
    import xml.etree.ElementTree as ET

    try:
        with zipfile.ZipFile(path) as bundle:
            payload = bundle.read("content.xml")
    except (KeyError, zipfile.BadZipFile) as error:
        raise ParseError(f"无法读取 ODT：{error}") from error
    try:
        root = ET.fromstring(payload)
    except ET.ParseError as error:
        raise ParseError(f"无法解析 ODT 内容：{error}") from error

    def local(tag: str) -> str:
        return tag.rsplit("}", 1)[-1]

    blocks: list[Block] = []

    def walk(node, *, in_header: bool = False) -> None:  # noqa: ANN001
        tag = local(node.tag)
        if tag in ("h", "p"):
            text = _odf_text(node)
            if not text:
                return
            if tag == "h":
                level = int(node.attrib.get(
                    "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}outline-level", "1") or 1)
                blocks.append(Block(kind=BLOCK_HEADING, text=text, level=max(1, min(6, level))))
            else:
                blocks.append(Block(kind=BLOCK_PARAGRAPH, text=text))
            return
        if tag == "table":
            rows: list[list[str]] = []
            for row in node:
                if local(row.tag) != "table-row":
                    continue
                cells = [_odf_text(cell) for cell in row if local(cell.tag) in ("table-cell", "covered-table-cell")]
                if any(cells):
                    rows.append(cells)
            if rows:
                blocks.append(Block(kind=BLOCK_TABLE, table=TableData(rows=rows, header=True)))
            return
        for child in node:
            walk(child, in_header=in_header)

    for child in root:
        if local(child.tag) in ("body", "document"):
            walk(child)
    if not blocks:
        raise ParseError("ODT 里没有可读正文")
    ir = DocumentIR(blocks=blocks, title=path.stem, format_key="odt")
    return ir


def _odf_text(node) -> str:  # noqa: ANN001
    parts: list[str] = []
    for element in node.iter():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag in ("p", "span", "a", "s"):
            if element.text:
                parts.append(element.text)
        if tag == "s":
            parts.append(" ")
        if element.tail and tag in ("span", "a"):
            parts.append(element.tail)
    text = "".join(parts)
    return re.sub(r"[ \t\u00a0]+", " ", text).strip()


def _parse_rtf(path: Path) -> DocumentIR:
    from modu_workbench.core.convert import office_io

    text = office_io._extract_rtf(path)      # noqa: SLF001  复用转换板块的轻量 RTF 解析
    if not text.strip():
        raise ParseError("这个 RTF 解析不出正文，建议先用 LibreOffice 另存为 DOCX")
    ir = parse_plain_text(text, title=path.stem)
    ir.format_key = "rtf"
    return ir


# ---------------------------------------------------------------- PDF


def _parse_pdf(path: Path, *, ocr: bool) -> DocumentIR:
    try:
        from pypdf import PdfReader
    except ImportError as error:  # pragma: no cover - pypdf 是随包依赖
        raise ParseError("缺少 pypdf，无法读取 PDF") from error

    try:
        reader = PdfReader(str(path))
    except Exception as error:  # noqa: BLE001
        raise ParseError(f"无法读取 PDF：{error}") from error

    blocks: list[Block] = []
    empty_pages = 0
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:  # noqa: BLE001  单页失败不影响其余页
            text = ""
        if not text:
            empty_pages += 1
            continue
        for chunk in re.split(r"\n\s*\n", text):
            chunk = chunk.strip()
            if not chunk:
                continue
            level = guess_heading_level(chunk.split("\n")[0])
            blocks.append(Block(
                kind=BLOCK_HEADING if level else BLOCK_PARAGRAPH,
                text=join_wrapped_lines(chunk.split("\n")) if not level else chunk,
                level=level,
                meta={"page": number},
            ))

    ir = DocumentIR(blocks=blocks, title=path.stem, format_key="pdf")
    ir.metadata["pages"] = len(reader.pages)
    ir.metadata["empty_pages"] = empty_pages
    ir.metadata["scanned"] = bool(empty_pages and not blocks)

    if empty_pages and not blocks:
        ir.warnings.append(
            "这份 PDF 抽不到文字（多为扫描件）：请用「OCR」功能识别，或先在墨软转换里转成图片再 OCR"
        )
    elif empty_pages:
        ir.warnings.append(f"有 {empty_pages} 页没有文字层（扫描页），可按需 OCR")

    if ocr and ir.metadata["scanned"]:
        ir = _ocr_pdf(path, ir)
    return ir


def _ocr_pdf(path: Path, ir: DocumentIR) -> DocumentIR:
    """扫描件 OCR：需要 PyMuPDF 或 pdf2image + tesseract；缺失时给出明确提示。"""
    from .ocr import ocr_pdf_pages

    pages = ocr_pdf_pages(path)
    if not pages:
        ir.warnings.append(
            "OCR 不可用：请安装 tesseract（或用 MODU_TESSERACT 指定）并安装 PyMuPDF，"
            "也可以先把 PDF 转成图片再 OCR"
        )
        return ir
    blocks: list[Block] = []
    for number, text in pages:
        for chunk in re.split(r"\n\s*\n", text):
            if chunk.strip():
                blocks.append(Block(kind=BLOCK_PARAGRAPH, text=chunk.strip(), meta={"page": number}))
    if blocks:
        ir.blocks = blocks
        ir.metadata["ocr"] = True
        ir.warnings = [warn for warn in ir.warnings if "OCR" not in warn]
        ir.warnings.append(f"已对 {len(pages)} 页执行 OCR（结果可能有识别误差，请校对）")
    return ir


# ---------------------------------------------------------------- 演示族


def _parse_slides_family(path: Path, format_key: str) -> DocumentIR:
    suffix = path.suffix.lower()
    if format_key == "pptx" or suffix in (".pptx", ".pptm"):
        return _parse_pptx(path)
    if format_key == "odp" or suffix == ".odp":
        try:
            return _parse_odp(path)
        except ParseError:
            pass
    if format_key == "dps" or suffix in (".dps", ".dpt"):
        container = sniff_ooxml(path)
        if container == "pptx":
            return _parse_pptx(path)
        if container == "docx":
            return _parse_docx(path)
    converted = _soffice_convert(path, "pptx")
    if converted is not None:
        try:
            ir = _parse_pptx(converted)
            ir.title = path.stem
            ir.format_key = format_key
            return ir
        finally:
            converted.unlink(missing_ok=True)
    converted = _soffice_convert(path, "txt")
    if converted is not None:
        try:
            ir = parse_plain_text(_read_text(converted), title=path.stem)
            ir.format_key = format_key
            ir.warnings.append("通过 LibreOffice 抽取纯文本：版式与图片未保留")
            return ir
        finally:
            converted.unlink(missing_ok=True)
    raise ParseError(
        f"{fmt.format_label(format_key)} 需要 LibreOffice（soffice）才能读取："
        "请安装后重试，或先另存为 PPTX。"
    )


_PPTX_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _parse_pptx(path: Path) -> DocumentIR:
    """PPTX：直接从 OOXML 取每页文字与备注（纯标准库，不依赖 python-pptx）。"""
    import xml.etree.ElementTree as ET

    try:
        bundle = zipfile.ZipFile(path)
    except (zipfile.BadZipFile, OSError) as error:
        raise ParseError(f"无法读取演示文稿：{error}") from error

    with bundle:
        slide_names = sorted(
            (name for name in bundle.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", name)),
            key=lambda item: int(re.findall(r"\d+", item)[-1]),
        )
        if not slide_names:
            raise ParseError("演示文稿里没有找到幻灯片")
        notes: dict[int, str] = {}
        for name in bundle.namelist():
            match = re.match(r"ppt/notesSlides/notesSlide(\d+)\.xml$", name)
            if match:
                notes[int(match.group(1))] = _pptx_texts(bundle.read(name), ET)
        blocks: list[Block] = []
        for index, name in enumerate(slide_names, start=1):
            texts = _pptx_texts(bundle.read(name), ET)
            if not texts and not notes.get(index):
                continue
            blocks.append(Block(
                kind=BLOCK_SLIDE,
                text="\n".join(texts).strip(),
                level=1,
                meta={"slide": index, "notes": notes.get(index, "")},
            ))
    if not blocks:
        raise ParseError("演示文稿里没有可读文字（可能整份都是图片）")
    ir = DocumentIR(blocks=blocks, title=path.stem, format_key="pptx")
    ir.metadata["slides"] = len(blocks)
    return ir


def _pptx_texts(payload: bytes, et_module) -> list[str]:  # noqa: ANN001
    try:
        root = et_module.fromstring(payload)
    except et_module.ParseError:
        return []
    texts: list[str] = []
    for paragraph in root.iter(f"{_PPTX_NS}p"):
        runs = [node.text or "" for node in paragraph.iter(f"{_PPTX_NS}t")]
        line = "".join(runs).strip()
        if line:
            texts.append(line)
    return texts


def _parse_odp(path: Path) -> DocumentIR:
    import xml.etree.ElementTree as ET

    try:
        with zipfile.ZipFile(path) as bundle:
            payload = bundle.read("content.xml")
    except (KeyError, zipfile.BadZipFile) as error:
        raise ParseError(f"无法读取 ODP：{error}") from error
    try:
        root = ET.fromstring(payload)
    except ET.ParseError as error:
        raise ParseError(f"无法解析 ODP 内容：{error}") from error
    blocks: list[Block] = []
    index = 0
    for node in root.iter():
        if node.tag.rsplit("}", 1)[-1] != "page":
            continue
        index += 1
        lines: list[str] = []
        for element in node.iter():
            if element.tag.rsplit("}", 1)[-1] == "p":
                text = _odf_text(element)
                if text:
                    lines.append(text)
        if lines:
            blocks.append(Block(kind=BLOCK_SLIDE, text="\n".join(lines), level=1,
                                meta={"slide": index}))
    if not blocks:
        raise ParseError("演示文稿里没有可读文字")
    return DocumentIR(blocks=blocks, title=path.stem, format_key="odp")


# ---------------------------------------------------------------- 数据族


def _parse_data_family(path: Path, format_key: str) -> DocumentIR:
    text = _read_text(path)
    pretty, structure_note = format_structured(text, format_key)
    ir = DocumentIR(
        title=path.stem,
        format_key=format_key,
        blocks=[Block(kind=BLOCK_CODE, text=pretty, style={"language": format_key})],
        metadata={"structure": structure_note},
    )
    return ir


def format_structured(text: str, format_key: str) -> tuple[str, str]:
    """把 JSON/XML/YAML/INI 格式化并给出结构摘要（返回 格式化文本, 摘要）。"""
    if format_key == "json":
        from modu_workbench.core.convert.text_io import JsonFormatError, json_pretty

        try:
            pretty = json_pretty(text, indent=2)
        except JsonFormatError as error:
            raise ParseError(str(error)) from error
        try:
            data = json.loads(text.lstrip("\ufeff") or "null")
        except json.JSONDecodeError:
            data = None
        if isinstance(data, dict):
            note = f"对象，{len(data)} 个键：" + "、".join(list(data)[:8])
        elif isinstance(data, list):
            note = f"数组，{len(data)} 项"
        else:
            note = "标量值"
        return pretty, note
    if format_key == "xml":
        import xml.dom.minidom as minidom

        try:
            document = minidom.parseString(text.encode("utf-8"))
        except Exception as error:  # noqa: BLE001
            raise ParseError(f"XML 语法错误：{error}") from error
        pretty = document.toprettyxml(indent="  ")
        pretty = "\n".join(line for line in pretty.splitlines() if line.strip())
        return pretty, f"根元素 <{document.documentElement.tagName}>"
    if format_key == "yaml":
        import yaml

        try:
            data = yaml.safe_load(text) or {}
        except yaml.YAMLError as error:
            raise ParseError(f"YAML 语法错误：{error}") from error
        pretty = yaml.safe_dump(data, allow_unicode=True, sort_keys=False, default_flow_style=False)
        if isinstance(data, dict):
            note = f"映射，{len(data)} 个键"
        elif isinstance(data, list):
            note = f"序列，{len(data)} 项"
        else:
            note = "标量值"
        return pretty.strip(), note
    # ini / toml / cfg
    import configparser

    parser = configparser.ConfigParser()
    parse_target = text if format_key != "toml" else _toml_to_ini(text)
    try:
        parser.read_string(parse_target)
    except configparser.Error as error:
        raise ParseError(f"INI 解析失败：{error}") from error
    buffer = io.StringIO()
    parser.write(buffer)
    sections = parser.sections()
    note = f"{len(sections)} 个分组：" + "、".join(sections[:8])
    return buffer.getvalue().strip(), note


def _toml_to_ini(text: str) -> str:
    """极简 TOML → INI 视图（只为结构化查看，不追求完整语义）。"""
    lines: list[str] = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        if line.startswith("[") and line.endswith("]"):
            lines.append(line)
            continue
        if "=" in line:
            key, _, value = line.partition("=")
            lines.append(f"{key.strip()} = {value.strip()}")
    return "\n".join(lines)


# ---------------------------------------------------------------- 图片


def _parse_image(path: Path, format_key: str, *, ocr: bool) -> DocumentIR:
    try:
        from PIL import Image
    except ImportError as error:  # pragma: no cover
        raise ParseError("缺少 Pillow，无法读取图片") from error

    ir = DocumentIR(title=path.stem, format_key=format_key)
    try:
        with Image.open(path) as image:
            width, height = image.size
            ir.metadata.update({
                "width": width, "height": height,
                "mode": image.mode, "frames": int(getattr(image, "n_frames", 1)),
            })
    except Exception as error:  # noqa: BLE001
        raise ParseError(f"无法读取图片：{error}") from error

    block = Block(kind=BLOCK_IMAGE, text=path.name, meta={"src": str(path)})
    ir.blocks = [block]
    if ocr:
        from .ocr import ocr_image

        text = ocr_image(path)
        if text:
            ir.blocks.append(Block(kind=BLOCK_PARAGRAPH, text=text,
                                   meta={"ocr": True, "source": str(path)}))
            ir.blocks[-1].ai_generated = False
            ir.metadata["ocr"] = True
            ir.warnings.append("OCR 结果可能有识别误差，请校对后再使用")
        else:
            ir.warnings.append(
                "OCR 不可用：请安装 tesseract 并用 MODU_TESSERACT 指定路径（识别语言默认 chi_sim+eng）"
            )
    return ir


# ---------------------------------------------------------------- EPUB


def _parse_epub(path: Path) -> DocumentIR:
    try:
        import ebooklib
        from ebooklib import epub
    except ImportError as error:  # pragma: no cover
        raise ParseError("缺少 ebooklib，无法读取 EPUB") from error

    try:
        book = epub.read_epub(str(path))
    except Exception as error:  # noqa: BLE001
        raise ParseError(f"无法读取 EPUB：{error}") from error

    blocks: list[Block] = []
    titles: list[str] = []
    for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
        try:
            html = item.get_content().decode("utf-8", "ignore")
        except Exception:  # noqa: BLE001
            continue
        chapter = parse_html(html)
        if chapter.blocks and not titles:
            first = chapter.blocks[0]
            if first.is_heading:
                titles.append(first.text)
        for block in chapter.blocks:
            if block.kind == BLOCK_HEADING and block.level > 2:
                block.level = 2
            blocks.append(block)
    if not blocks:
        raise ParseError("EPUB 里没有可读正文")
    title = path.stem
    try:
        meta_title = book.get_metadata("DC", "title")
        if meta_title:
            title = str(meta_title[0][0]) or title
    except Exception:  # noqa: BLE001
        pass
    return DocumentIR(blocks=blocks, title=title, format_key="epub")


# ---------------------------------------------------------------- LibreOffice


def _soffice_convert(path: Path, target: str) -> Optional[Path]:
    from modu_workbench.core.convert import office_io

    try:
        return office_io.soffice_convert(path, target=target)
    except Exception:  # noqa: BLE001  没有 soffice → None，由调用方给中文提示
        return None


def soffice_available() -> bool:
    from modu_workbench.core.convert.office_io import find_soffice

    return bool(find_soffice())


def ocr_available() -> bool:
    from .ocr import available

    return available()


def describe_support() -> list[tuple[str, str]]:
    """(能力, 是否可用) 列表 —— 设置页展示"本机能不能读 ODS/PPT/OCR"。"""
    return [
        ("LibreOffice（ODS / 旧版 DOC·PPT / WPS 旧格式）", soffice_available()),
        ("OCR（扫描件与图片识别）", ocr_available()),
    ]


__all__ = [
    "MAX_TABLE_ROWS",
    "ParseError",
    "describe_support",
    "format_structured",
    "guess_heading_level",
    "join_wrapped_lines",
    "ocr_available",
    "parse_document",
    "parse_html",
    "parse_markdown",
    "parse_plain_text",
    "read_document_text",
    "sniff_ooxml",
    "soffice_available",
    "table_from_rows",
]
